#!/usr/bin/env python3
"""
Create Master Template - functional template with working formulas for any asset type
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

YELLOW_FILL = PatternFill(start_color="FFFFC7", end_color="FFFFC7", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
NAVY_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
LIGHT_GRAY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

BLUE_BOLD = Font(color="0000FF", bold=True)
BLACK_FONT = Font(color="000000")
BLACK_BOLD = Font(color="000000", bold=True)
WHITE_BOLD = Font(color="FFFFFF", bold=True)

THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

CURRENCY_FORMAT = '#,##0.0_);(#,##0.0)'
PERCENT_FORMAT = '0.0%'
MULTIPLE_FORMAT = '0.00"x"'
INTEGER_FORMAT = '#,##0'

def col_letter(col_num):
    return get_column_letter(col_num)

def apply_input_style(ws, row, col, value=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    cell.fill = YELLOW_FILL
    cell.font = BLUE_BOLD
    cell.border = THIN_BORDER
    return cell

def apply_calc_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = PatternFill(fill_type=None)
    cell.font = BLACK_FONT
    cell.border = THIN_BORDER
    return cell

def apply_output_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = GREEN_FILL
    cell.font = BLACK_BOLD
    cell.border = THIN_BORDER
    return cell

def apply_header_style(ws, row, col, text):
    cell = ws.cell(row=row, column=col)
    cell.value = text
    cell.fill = NAVY_FILL
    cell.font = WHITE_BOLD
    cell.alignment = Alignment(horizontal='left')
    return cell

def apply_warning_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = ORANGE_FILL
    cell.font = BLACK_FONT
    cell.border = THIN_BORDER
    return cell

def set_label(ws, row, label, units="", notes=""):
    ws.cell(row=row, column=1, value=label).font = BLACK_FONT
    ws.cell(row=row, column=2, value=units).font = Font(color="666666", italic=True)
    ws.cell(row=row, column=3, value=notes).font = Font(color="666666", italic=True)

def create_master_template():
    """Create master template with working formulas"""

    wb = openpyxl.Workbook()

    # ========================================================================
    # INSTRUCTIONS SHEET
    # ========================================================================
    ws_inst = wb.active
    ws_inst.title = "Instructions"

    ws_inst.column_dimensions['A'].width = 20
    ws_inst.column_dimensions['B'].width = 80

    instructions = [
        ("MASTER TEMPLATE INSTRUCTIONS", ""),
        ("", ""),
        ("PURPOSE", "This template provides a functional framework for building infrastructure project finance models. All formulas are working and can be adapted for any asset type."),
        ("", ""),
        ("ASSET TYPES SUPPORTED", ""),
        ("• Thermal Generation", "CCGT, Peaker, Coal (modify fuel cost calculations)"),
        ("• Renewables", "Solar, Wind, BESS (remove fuel cost, add degradation)"),
        ("• Transmission", "Add construction period, use RAB-based returns"),
        ("• Midstream", "Volume-based revenue, growth-decline profile"),
        ("", ""),
        ("DEBT STRUCTURES", ""),
        ("• Cash Sweep", "For merchant assets with volatile CFADS"),
        ("• Sculpted", "For contracted assets with predictable CFADS"),
        ("", ""),
        ("KEY MODIFICATIONS BY ASSET", ""),
        ("CCGT/Peaker", "Use provided fuel cost formula: Gen × HR × Gas / 1e9"),
        ("Solar/Wind", "Remove fuel cost row, add degradation to generation"),
        ("Transmission", "Add construction period rows Y0-Y2, use COD timing"),
        ("Midstream", "Use volume × fee revenue, add growth/decline formula"),
        ("", ""),
        ("FORMULA CONVENTIONS", ""),
        ("$D$ references", "Inputs in column D are typically fixed references"),
        ("Column references", "E through N represent Years 1-10"),
        ("Escalation", "All prices use (1+esc%)^(year-1) pattern"),
        ("", ""),
        ("COMMON ERRORS TO AVOID", ""),
        ("1. Fuel cost divisor", "Must divide by 1,000,000,000 (not 1,000,000)"),
        ("2. Peaker generation", "Use dispatch hours, NOT capacity factor × 8,760"),
        ("3. Capacity revenue", "Multiply by 1,000 (capacity price is per kW, not MW)"),
        ("4. DSRA target", "Reference NEXT year's debt service, not current"),
    ]

    row = 1
    for label, explanation in instructions:
        if label in ["MASTER TEMPLATE INSTRUCTIONS", "PURPOSE", "ASSET TYPES SUPPORTED",
                     "DEBT STRUCTURES", "KEY MODIFICATIONS BY ASSET", "FORMULA CONVENTIONS",
                     "COMMON ERRORS TO AVOID"]:
            ws_inst.cell(row=row, column=1, value=label).font = Font(bold=True, size=12)
            ws_inst.cell(row=row, column=1).fill = NAVY_FILL
            ws_inst.cell(row=row, column=1).font = WHITE_BOLD
            ws_inst.cell(row=row, column=2).fill = NAVY_FILL
        else:
            ws_inst.cell(row=row, column=1, value=label).font = BLACK_BOLD
            ws_inst.cell(row=row, column=2, value=explanation).font = BLACK_FONT
        ws_inst.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        row += 1

    # ========================================================================
    # MODEL SHEET (Generic Template)
    # ========================================================================
    ws = wb.create_sheet("Model")
    rows = {}

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 38
    for col in range(4, 15):
        ws.column_dimensions[col_letter(col)].width = 14

    # Year Headers
    row = 1
    rows['year_header'] = row
    ws.cell(row=row, column=1, value="Infrastructure Model Template").font = Font(bold=True, size=14)
    ws.cell(row=row, column=4, value=0)
    for yr in range(1, 11):
        ws.cell(row=row, column=4+yr, value=yr)
    for col in range(4, 15):
        cell = ws.cell(row=row, column=col)
        cell.font = BLACK_BOLD
        cell.alignment = Alignment(horizontal='center')

    # Audit Strip
    row = 2
    apply_header_style(ws, row, 1, "═══ AUDIT STRIP ═══")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 3; rows['audit_entry_ev'] = row
    set_label(ws, row, "Entry EV", "$mm")

    row = 4; rows['audit_y1_ebitda'] = row
    set_label(ws, row, "Y1 EBITDA", "$mm")

    row = 5; rows['audit_min_dscr'] = row
    set_label(ws, row, "Min DSCR", "x")

    row = 6; rows['audit_levered_irr'] = row
    set_label(ws, row, "Levered IRR", "%")

    row = 7; rows['audit_moic'] = row
    set_label(ws, row, "MOIC", "x")

    row = 8
    apply_header_style(ws, row, 1, "═══ QA CHECKS ═══")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 9; rows['qa_su_balance'] = row
    set_label(ws, row, "S&U Balance")

    row = 10; rows['qa_debt_payoff'] = row
    set_label(ws, row, "Debt Payoff at Exit")

    row = 12
    # Transaction Inputs
    apply_header_style(ws, row, 1, "TRANSACTION INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 13; rows['entry_ev'] = row
    set_label(ws, row, "Entry EV", "$mm", "[INPUT] Purchase price")
    apply_input_style(ws, row, 4, 100)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 14; rows['txn_fees'] = row
    set_label(ws, row, "Transaction Fees", "%", "[INPUT] Legal, advisory")
    apply_input_style(ws, row, 4, 0.015)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 15; rows['exit_multiple'] = row
    set_label(ws, row, "Exit Multiple", "x", "[INPUT] On exit EBITDA")
    apply_input_style(ws, row, 4, 7.0)
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    row = 16; rows['exit_year'] = row
    set_label(ws, row, "Exit Year", "#", "[INPUT] Hold period")
    apply_input_style(ws, row, 4, 7)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 18
    # Asset Inputs
    apply_header_style(ws, row, 1, "ASSET INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 19; rows['capacity'] = row
    set_label(ws, row, "Capacity", "MW/units", "[INPUT] Asset size")
    apply_input_style(ws, row, 4, 100)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 20; rows['utilization'] = row
    set_label(ws, row, "Utilization/CF", "%", "[INPUT] Capacity factor or utilization")
    apply_input_style(ws, row, 4, 0.50)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 21; rows['escalation'] = row
    set_label(ws, row, "Escalation", "%/yr", "[INPUT] All prices/costs")
    apply_input_style(ws, row, 4, 0.025)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 23
    # Revenue Inputs
    apply_header_style(ws, row, 1, "REVENUE INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 24; rows['price_y1'] = row
    set_label(ws, row, "Price Y1", "$/unit", "[INPUT] Base year price")
    apply_input_style(ws, row, 4, 50)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 26
    # Cost Inputs
    apply_header_style(ws, row, 1, "COST INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 27; rows['opex_y1'] = row
    set_label(ws, row, "Operating Costs Y1", "$mm", "[INPUT] Base year opex")
    apply_input_style(ws, row, 4, 5)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 29
    # Financing Inputs
    apply_header_style(ws, row, 1, "FINANCING INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 30; rows['leverage'] = row
    set_label(ws, row, "Leverage", "%", "[INPUT] Debt / EV")
    apply_input_style(ws, row, 4, 0.40)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 31; rows['interest_rate'] = row
    set_label(ws, row, "Interest Rate", "%", "[INPUT] All-in rate")
    apply_input_style(ws, row, 4, 0.06)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 32; rows['debt_tenor'] = row
    set_label(ws, row, "Debt Tenor", "years", "[INPUT] Amortization period")
    apply_input_style(ws, row, 4, 7)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 33; rows['dsra_months'] = row
    set_label(ws, row, "DSRA Months", "months", "[INPUT] Reserve sizing")
    apply_input_style(ws, row, 4, 6)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 34; rows['sweep_pct'] = row
    set_label(ws, row, "Cash Sweep %", "%", "[INPUT] For sweep structure")
    apply_input_style(ws, row, 4, 0.75)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 35; rows['lockup_dscr'] = row
    set_label(ws, row, "Lock-up DSCR", "x", "[INPUT] Distribution threshold")
    apply_input_style(ws, row, 4, 1.10)
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    row = 37
    # Tax Inputs
    apply_header_style(ws, row, 1, "TAX INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 38; rows['tax_rate'] = row
    set_label(ws, row, "Tax Rate", "%", "[INPUT] Blended rate")
    apply_input_style(ws, row, 4, 0.25)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 39; rows['dep_basis'] = row
    set_label(ws, row, "Depreciable Basis", "%", "[INPUT] % of EV")
    apply_input_style(ws, row, 4, 0.85)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 40; rows['dep_life'] = row
    set_label(ws, row, "Depreciation Life", "years", "[INPUT] SL life")
    apply_input_style(ws, row, 4, 15)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 42
    # Calculated Items
    apply_header_style(ws, row, 1, "CALCULATED ITEMS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 43; rows['debt_amount'] = row
    set_label(ws, row, "Debt Amount", "$mm", "=EV × Leverage")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['leverage']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 44; rows['annual_dep'] = row
    set_label(ws, row, "Annual Depreciation", "$mm", "=EV × Basis / Life")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['dep_basis']}/$D${rows['dep_life']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 45; rows['sched_prin'] = row
    set_label(ws, row, "Scheduled Principal", "$mm/yr", "=Debt / Tenor")
    apply_calc_style(ws, row, 4, f"=$D${rows['debt_amount']}/$D${rows['debt_tenor']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 47
    # Operating Model
    apply_header_style(ws, row, 1, "OPERATING MODEL")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 48; rows['price'] = row
    set_label(ws, row, "Price", "$/unit", "Escalated")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['price_y1']}*(1+$D${rows['escalation']})^({yr}-1)"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 49; rows['revenue'] = row
    set_label(ws, row, "Revenue", "$mm", "[Customize formula]")
    for col in range(5, 15):
        # Generic: Capacity × Utilization × Price × Hours / conversion
        formula = f"=$D${rows['capacity']}*$D${rows['utilization']}*{col_letter(col)}{rows['price']}*8760/1000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 50; rows['opex'] = row
    set_label(ws, row, "Operating Costs", "$mm", "Escalated")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['opex_y1']}*(1+$D${rows['escalation']})^({yr}-1)"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 51; rows['ebitda'] = row
    set_label(ws, row, "EBITDA", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['revenue']}-{col_letter(col)}{rows['opex']}"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 53
    # Cash Flow
    apply_header_style(ws, row, 1, "CASH FLOW")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 54; rows['ebit'] = row
    set_label(ws, row, "EBIT", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['ebitda']}-$D${rows['annual_dep']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 55; rows['taxes'] = row
    set_label(ws, row, "Taxes", "$mm")
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['ebit']})*$D${rows['tax_rate']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 56; rows['cfads'] = row
    set_label(ws, row, "CFADS", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['ebitda']}-{col_letter(col)}{rows['taxes']}"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 58
    # Debt Schedule
    apply_header_style(ws, row, 1, "DEBT SCHEDULE")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 59; rows['debt_beg'] = row
    set_label(ws, row, "Beginning Balance", "$mm")
    apply_calc_style(ws, row, 5, f"=$D${rows['debt_amount']}")
    ws.cell(row=row, column=5).number_format = CURRENCY_FORMAT

    row = 60; rows['interest'] = row
    set_label(ws, row, "Interest", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['debt_beg']}*$D${rows['interest_rate']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 61; rows['prin_sched'] = row
    set_label(ws, row, "Scheduled Principal", "$mm")
    for col in range(5, 15):
        formula = f"=MIN($D${rows['sched_prin']},{col_letter(col)}{rows['debt_beg']})"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 62; rows['ds_pre'] = row
    set_label(ws, row, "Debt Service (Pre-Sweep)", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['interest']}+{col_letter(col)}{rows['prin_sched']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 63; rows['excess_cash'] = row
    set_label(ws, row, "Excess Cash", "$mm")

    row = 64; rows['sweep'] = row
    set_label(ws, row, "Cash Sweep", "$mm")
    for col in range(5, 15):
        formula = f"=MIN(MAX(0,{col_letter(col)}{rows['excess_cash']})*$D${rows['sweep_pct']},MAX(0,{col_letter(col)}{rows['debt_beg']}-{col_letter(col)}{rows['prin_sched']}))"
        apply_warning_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 65; rows['total_prin'] = row
    set_label(ws, row, "Total Principal", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['prin_sched']}+{col_letter(col)}{rows['sweep']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 66; rows['debt_end'] = row
    set_label(ws, row, "Ending Balance", "$mm")
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['debt_beg']}-{col_letter(col)}{rows['total_prin']})"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    # Fix beg balance references
    for col in range(6, 15):
        prev_col = col_letter(col - 1)
        formula = f"={prev_col}{rows['debt_end']}"
        apply_calc_style(ws, rows['debt_beg'], col, formula)

    row = 67; rows['dscr'] = row
    set_label(ws, row, "DSCR", "x")
    for col in range(5, 15):
        formula = f"=IF({col_letter(col)}{rows['ds_pre']}>0,{col_letter(col)}{rows['cfads']}/{col_letter(col)}{rows['ds_pre']},0)"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = MULTIPLE_FORMAT

    row = 69
    # DSRA
    apply_header_style(ws, row, 1, "DSRA")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 70; rows['dsra_target'] = row
    set_label(ws, row, "DSRA Target", "$mm")
    apply_calc_style(ws, row, 4, f"=E{rows['ds_pre']}*$D${rows['dsra_months']}/12")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 14):
        next_col = col_letter(col + 1)
        formula = f"={next_col}{rows['ds_pre']}*$D${rows['dsra_months']}/12"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
    apply_calc_style(ws, row, 14, "=0")
    ws.cell(row=row, column=14).number_format = CURRENCY_FORMAT

    row = 71; rows['dsra_beg'] = row
    set_label(ws, row, "DSRA Beginning", "$mm")
    apply_calc_style(ws, row, 4, "=0")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 72; rows['dsra_funding'] = row
    set_label(ws, row, "DSRA Funding/(Release)", "$mm")
    for col in range(4, 15):
        formula = f"={col_letter(col)}{rows['dsra_target']}-{col_letter(col)}{rows['dsra_beg']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 73; rows['dsra_end'] = row
    set_label(ws, row, "DSRA Ending", "$mm")
    for col in range(4, 15):
        formula = f"={col_letter(col)}{rows['dsra_beg']}+{col_letter(col)}{rows['dsra_funding']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    # Fix DSRA beg references
    for col in range(5, 15):
        prev_col = col_letter(col - 1)
        formula = f"={prev_col}{rows['dsra_end']}"
        apply_calc_style(ws, rows['dsra_beg'], col, formula)

    # Add excess cash formulas
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['cfads']}-{col_letter(col)}{rows['ds_pre']}-{col_letter(col)}{rows['dsra_funding']})"
        apply_calc_style(ws, rows['excess_cash'], col, formula)
        ws.cell(row=rows['excess_cash'], column=col).number_format = CURRENCY_FORMAT

    row = 75
    # Distribution Waterfall
    apply_header_style(ws, row, 1, "DISTRIBUTION WATERFALL")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 76; rows['cash_avail'] = row
    set_label(ws, row, "Cash Available", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['cfads']}-({col_letter(col)}{rows['interest']}+{col_letter(col)}{rows['total_prin']})-{col_letter(col)}{rows['dsra_funding']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 77; rows['lockup_test'] = row
    set_label(ws, row, "Lock-up Test", "")
    for col in range(5, 15):
        formula = f'=IF({col_letter(col)}{rows["dscr"]}>=$D${rows["lockup_dscr"]},"PASS","LOCKED")'
        apply_calc_style(ws, row, col, formula)

    row = 78; rows['distributable'] = row
    set_label(ws, row, "Distributable Cash", "$mm")
    for col in range(5, 15):
        formula = f'=IF({col_letter(col)}{rows["lockup_test"]}="PASS",MAX(0,{col_letter(col)}{rows["cash_avail"]}),0)'
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 80
    # Sources & Uses
    apply_header_style(ws, row, 1, "SOURCES & USES")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 81; rows['use_purchase'] = row
    set_label(ws, row, "Purchase Price", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 82; rows['use_fees'] = row
    set_label(ws, row, "Transaction Fees", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['txn_fees']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 83; rows['use_dsra'] = row
    set_label(ws, row, "DSRA Funding", "$mm")
    apply_calc_style(ws, row, 4, f"=D{rows['dsra_target']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 84; rows['total_uses'] = row
    set_label(ws, row, "Total Uses", "$mm")
    apply_output_style(ws, row, 4, f"=D{rows['use_purchase']}+D{rows['use_fees']}+D{rows['use_dsra']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 86; rows['source_debt'] = row
    set_label(ws, row, "Debt", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['debt_amount']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 87; rows['source_equity'] = row
    set_label(ws, row, "Equity", "$mm")
    apply_calc_style(ws, row, 4, f"=D{rows['total_uses']}-D{rows['source_debt']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 88; rows['total_sources'] = row
    set_label(ws, row, "Total Sources", "$mm")
    apply_output_style(ws, row, 4, f"=D{rows['source_debt']}+D{rows['source_equity']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 89; rows['su_check'] = row
    set_label(ws, row, "Balance Check", "")
    apply_output_style(ws, row, 4, f'=IF(ABS(D{rows["total_uses"]}-D{rows["total_sources"]})<0.01,"PASS","FAIL")')

    row = 91
    # Exit & Returns
    apply_header_style(ws, row, 1, "EXIT & RETURNS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 92; rows['exit_ev'] = row
    set_label(ws, row, "Exit EV", "$mm")
    apply_calc_style(ws, row, 11, f"=$D${rows['exit_multiple']}*K{rows['ebitda']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 93; rows['exit_debt'] = row
    set_label(ws, row, "Exit Debt Payoff", "$mm")
    apply_calc_style(ws, row, 11, f"=K{rows['debt_end']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 94; rows['exit_dsra'] = row
    set_label(ws, row, "DSRA Release", "$mm")
    apply_calc_style(ws, row, 11, f"=K{rows['dsra_end']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 95; rows['exit_equity'] = row
    set_label(ws, row, "Exit Equity Proceeds", "$mm")
    apply_output_style(ws, row, 11, f"=K{rows['exit_ev']}-K{rows['exit_debt']}+K{rows['exit_dsra']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 97; rows['equity_cf'] = row
    set_label(ws, row, "Equity Cash Flow", "$mm")
    apply_calc_style(ws, row, 4, f"=-D{rows['source_equity']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 11):
        formula = f"={col_letter(col)}{rows['distributable']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
    apply_calc_style(ws, row, 11, f"=K{rows['distributable']}+K{rows['exit_equity']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT
    for col in range(12, 15):
        apply_calc_style(ws, row, col, "=0")
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 98; rows['levered_irr'] = row
    set_label(ws, row, "Levered IRR", "%")
    apply_output_style(ws, row, 4, f"=IRR(D{rows['equity_cf']}:K{rows['equity_cf']})")
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 99; rows['moic'] = row
    set_label(ws, row, "MOIC", "x")
    apply_output_style(ws, row, 4, f"=SUMIF(D{rows['equity_cf']}:K{rows['equity_cf']},\">0\")/ABS(D{rows['equity_cf']})")
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    # Populate Audit Strip
    apply_output_style(ws, rows['audit_entry_ev'], 4, f"=$D${rows['entry_ev']}")
    ws.cell(row=rows['audit_entry_ev'], column=4).number_format = CURRENCY_FORMAT

    apply_output_style(ws, rows['audit_y1_ebitda'], 4, f"=E{rows['ebitda']}")
    ws.cell(row=rows['audit_y1_ebitda'], column=4).number_format = CURRENCY_FORMAT

    apply_output_style(ws, rows['audit_min_dscr'], 4, f"=MIN(E{rows['dscr']}:K{rows['dscr']})")
    ws.cell(row=rows['audit_min_dscr'], column=4).number_format = MULTIPLE_FORMAT

    apply_output_style(ws, rows['audit_levered_irr'], 4, f"=D{rows['levered_irr']}")
    ws.cell(row=rows['audit_levered_irr'], column=4).number_format = PERCENT_FORMAT

    apply_output_style(ws, rows['audit_moic'], 4, f"=D{rows['moic']}")
    ws.cell(row=rows['audit_moic'], column=4).number_format = MULTIPLE_FORMAT

    apply_output_style(ws, rows['qa_su_balance'], 4, f"=D{rows['su_check']}")
    apply_output_style(ws, rows['qa_debt_payoff'], 4, f'=IF(K{rows["debt_end"]}<1,"PASS","FAIL")')

    ws.freeze_panes = 'E2'

    wb.save('/mnt/user-data/outputs/Master_Template.xlsx')
    print('Created Master_Template.xlsx')

if __name__ == "__main__":
    create_master_template()
