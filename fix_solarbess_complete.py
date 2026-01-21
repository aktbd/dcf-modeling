#!/usr/bin/env python3
"""
Complete fix for Model_3_SolarBESS.xlsx
Key features: Solar degradation, BESS revenue, ITC, Sculpted debt
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'
CURRENCY_FMT = '#,##0.0_);(#,##0.0)'

def fix_solarbess():
    filepath = f"{DELIV_DIR}/Model_3_SolarBESS.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    # ===== INPUT ROW MAPPING =====
    inp = {
        'entry_ev': 27,       # 145
        'txn_fees': 28,       # 0.015
        'itc_rate': 29,       # 0.3
        'exit_mult': 30,      # 8
        'exit_year': 31,      # 10
        'ac_capacity': 34,    # 115
        'cap_factor': 35,     # 0.26
        'degradation': 36,    # 0.005
        'ppa_price': 39,      # 35
        'ppa_esc': 40,        # 0.015
        'bess_rev': 41,       # 4.5
        'bess_esc': 42,       # 0.02
        'om_cost': 45,        # 2
        'om_esc': 46,         # 0.02
        'aug_year': 47,       # 7
        'aug_cost': 48,       # 5
        'leverage': 51,       # 0.35
        'int_rate': 52,       # 0.055
        'debt_tenor': 53,     # 10
        'dsra_months': 54,    # 6
        'target_dscr': 55,    # 1.35
        'lockup': 56,         # 1.1
        'tax_rate': 59,       # 0.25
        'macrs_years': 60,    # 6
        'post_macrs_tax': 61, # 0.15
    }

    calc = {
        'y1_gen': 64,
        'debt_amount': 65,
        'itc_benefit': 66,
    }

    op = {
        'generation': 69,
        'ppa_price': 70,
        'ppa_rev': 72,
        'bess_rev': 73,
        'total_rev': 74,
        'om_cost': 76,
        'ebitda': 78,
        'taxes': 81,
        'augmentation': 82,
        'cfads': 83,
    }

    # Sculpted debt (different from cash sweep)
    debt = {
        'beg_bal': 86,
        'interest': 87,
        'target_ds': 88,
        'sculpted_prin': 89,
        'actual_ds': 90,
        'end_bal': 91,
        'dscr': 92,
    }

    dsra = {
        'target': 95,
        'beg': 96,
        'funding': 97,
        'end': 98,
    }

    dist = {
        'avail': 101,
        'lockup': 102,
        'distributable': 103,
    }

    su = {
        'purchase': 107,
        'txn_fees': 108,
        'dsra_init': 109,
        'total_uses': 110,
        'debt': 113,
        'pre_itc_equity': 114,
        'itc_benefit': 115,
        'net_equity': 116,
        'total_sources': 117,
        'balance': 118,
    }

    exit_r = {
        'exit_ev': 121,
        'exit_debt': 122,
        'exit_dsra': 123,
        'exit_equity': 124,
        'equity_cf': 127,
        'irr': 128,
        'moic': 129,
        'unlev_cf': 132,
        'unlev_irr': 133,
    }

    print("="*60)
    print("FIXING MODEL_3_SOLARBESS.xlsx")
    print("="*60)

    # ===== 1. CALCULATED ITEMS =====
    print("[1] Fixing Calculated Items...")

    # Y1 Generation = AC Capacity * CF * 8760
    ws.cell(row=calc['y1_gen'], column=4).value = f"=$D${inp['ac_capacity']}*$D${inp['cap_factor']}*8760"

    # Debt Amount = Entry EV * Leverage
    ws.cell(row=calc['debt_amount'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"

    # ITC Benefit = Entry EV * ITC Rate
    ws.cell(row=calc['itc_benefit'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['itc_rate']}"

    # ===== 2. OPERATING MODEL =====
    print("[2] Fixing Operating Model...")

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Solar Generation with degradation: Y1 Gen * (1 - Deg)^(Year-1)
        ws.cell(row=op['generation'], column=col).value = f"=$D${calc['y1_gen']}*(1-$D${inp['degradation']})^({year}-1)"

        # PPA Price with escalation
        ws.cell(row=op['ppa_price'], column=col).value = f"=$D${inp['ppa_price']}*(1+$D${inp['ppa_esc']})^({year}-1)"

        # PPA Revenue = Generation * PPA Price / 1E6
        ws.cell(row=op['ppa_rev'], column=col).value = f"={c}{op['generation']}*{c}{op['ppa_price']}/1000000"

        # BESS Revenue with escalation
        ws.cell(row=op['bess_rev'], column=col).value = f"=$D${inp['bess_rev']}*(1+$D${inp['bess_esc']})^({year}-1)"

        # Total Revenue
        ws.cell(row=op['total_rev'], column=col).value = f"={c}{op['ppa_rev']}+{c}{op['bess_rev']}"

        # O&M Cost with escalation
        ws.cell(row=op['om_cost'], column=col).value = f"=$D${inp['om_cost']}*(1+$D${inp['om_esc']})^({year}-1)"

        # EBITDA = Revenue - O&M
        ws.cell(row=op['ebitda'], column=col).value = f"={c}{op['total_rev']}-{c}{op['om_cost']}"
        ws.cell(row=op['ebitda'], column=col).fill = GREEN_FILL
        ws.cell(row=op['ebitda'], column=col).font = BLACK_BOLD
        ws.cell(row=op['ebitda'], column=col).number_format = CURRENCY_FMT

        # Taxes (simplified - use MACRS shield for first 6 years)
        ws.cell(row=op['taxes'], column=col).value = \
            f"=IF({year}<=$D${inp['macrs_years']},0,MAX(0,{c}{op['ebitda']})*$D${inp['post_macrs_tax']})"

        # BESS Augmentation (Year 7 only)
        ws.cell(row=op['augmentation'], column=col).value = f"=IF({year}=$D${inp['aug_year']},$D${inp['aug_cost']},0)"

        # CFADS = EBITDA - Taxes - Augmentation
        ws.cell(row=op['cfads'], column=col).value = f"={c}{op['ebitda']}-{c}{op['taxes']}-{c}{op['augmentation']}"
        ws.cell(row=op['cfads'], column=col).fill = GREEN_FILL
        ws.cell(row=op['cfads'], column=col).font = BLACK_BOLD

    # ===== 3. SCULPTED DEBT SCHEDULE =====
    print("[3] Fixing Sculpted Debt Schedule...")

    ws.cell(row=debt['beg_bal'], column=4).value = 0
    ws.cell(row=debt['end_bal'], column=4).value = f"=$D${calc['debt_amount']}"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)

        # Beginning Balance
        if year == 1:
            ws.cell(row=debt['beg_bal'], column=col).value = f"=$D${calc['debt_amount']}"
        else:
            ws.cell(row=debt['beg_bal'], column=col).value = f"={prev}{debt['end_bal']}"

        # Interest = Beg Bal * Interest Rate
        ws.cell(row=debt['interest'], column=col).value = f"={c}{debt['beg_bal']}*$D${inp['int_rate']}"

        # Target DS = CFADS / Target DSCR (sculpted)
        ws.cell(row=debt['target_ds'], column=col).value = f"={c}{op['cfads']}/$D${inp['target_dscr']}"

        # Sculpted Principal = MIN(MAX(0, Target DS - Interest), Beg Bal)
        ws.cell(row=debt['sculpted_prin'], column=col).value = \
            f"=MIN(MAX(0,{c}{debt['target_ds']}-{c}{debt['interest']}),{c}{debt['beg_bal']})"

        # Actual DS = Interest + Sculpted Principal
        ws.cell(row=debt['actual_ds'], column=col).value = f"={c}{debt['interest']}+{c}{debt['sculpted_prin']}"

        # Ending Balance = MAX(0, Beg Bal - Sculpted Principal)
        ws.cell(row=debt['end_bal'], column=col).value = f"=MAX(0,{c}{debt['beg_bal']}-{c}{debt['sculpted_prin']})"

        # Actual DSCR
        ws.cell(row=debt['dscr'], column=col).value = f"=IF({c}{debt['actual_ds']}>0,{c}{op['cfads']}/{c}{debt['actual_ds']},0)"
        ws.cell(row=debt['dscr'], column=col).number_format = MULTIPLE_FMT

    # ===== 4. DSRA =====
    print("[4] Fixing DSRA...")

    ws.cell(row=dsra['beg'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)
        next_c = get_column_letter(col + 1) if col < 14 else c

        # DSRA Target = Next Year's Actual DS * (DSRA Months / 12)
        if col < 14:
            ws.cell(row=dsra['target'], column=col).value = f"={next_c}{debt['actual_ds']}*$D${inp['dsra_months']}/12"
        else:
            ws.cell(row=dsra['target'], column=col).value = 0

        if year == 1:
            ws.cell(row=dsra['beg'], column=col).value = f"=D{su['dsra_init']}"
        else:
            ws.cell(row=dsra['beg'], column=col).value = f"={prev}{dsra['end']}"

        ws.cell(row=dsra['funding'], column=col).value = f"={c}{dsra['target']}-{c}{dsra['beg']}"
        ws.cell(row=dsra['end'], column=col).value = f"={c}{dsra['beg']}+{c}{dsra['funding']}"

    # ===== 5. DISTRIBUTION WATERFALL =====
    print("[5] Fixing Distribution Waterfall...")

    for col in range(5, 15):
        c = get_column_letter(col)

        ws.cell(row=dist['avail'], column=col).value = \
            f"={c}{op['cfads']}-{c}{debt['actual_ds']}-{c}{dsra['funding']}"
        ws.cell(row=dist['lockup'], column=col).value = f"=IF({c}{debt['dscr']}>=$D${inp['lockup']},\"PASS\",\"LOCKED\")"
        ws.cell(row=dist['distributable'], column=col).value = \
            f"=IF({c}{dist['lockup']}=\"PASS\",MAX(0,{c}{dist['avail']}),0)"

    # ===== 6. SOURCES & USES =====
    print("[6] Fixing Sources & Uses...")

    ws.cell(row=su['purchase'], column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=su['txn_fees'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['txn_fees']}"
    ws.cell(row=su['dsra_init'], column=4).value = f"=E{dsra['target']}"
    ws.cell(row=su['total_uses'], column=4).value = f"=D{su['purchase']}+D{su['txn_fees']}+D{su['dsra_init']}"

    ws.cell(row=su['debt'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
    ws.cell(row=su['pre_itc_equity'], column=4).value = f"=D{su['total_uses']}-D{su['debt']}"
    ws.cell(row=su['itc_benefit'], column=4).value = f"=$D${calc['itc_benefit']}"
    ws.cell(row=su['net_equity'], column=4).value = f"=D{su['pre_itc_equity']}-D{su['itc_benefit']}"
    ws.cell(row=su['total_sources'], column=4).value = f"=D{su['debt']}+D{su['net_equity']}+D{su['itc_benefit']}"
    ws.cell(row=su['balance'], column=4).value = f"=IF(ABS(D{su['total_uses']}-D{su['total_sources']})<0.01,\"PASS\",\"FAIL\")"

    # ===== 7. EXIT & RETURNS =====
    print("[7] Fixing Exit & Returns...")

    exit_year = 10
    exit_col = get_column_letter(4 + exit_year)  # N for Y10

    ws.cell(row=exit_r['exit_ev'], column=4).value = f"={exit_col}{op['ebitda']}*$D${inp['exit_mult']}"
    ws.cell(row=exit_r['exit_debt'], column=4).value = f"={exit_col}{debt['end_bal']}"
    ws.cell(row=exit_r['exit_dsra'], column=4).value = f"={exit_col}{dsra['end']}"
    ws.cell(row=exit_r['exit_equity'], column=4).value = f"=D{exit_r['exit_ev']}-D{exit_r['exit_debt']}+D{exit_r['exit_dsra']}"

    # Equity CF (use Net Equity for Y0)
    ws.cell(row=exit_r['equity_cf'], column=4).value = f"=-D{su['net_equity']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}"
        elif year == exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}+D{exit_r['exit_equity']}"
        else:
            ws.cell(row=exit_r['equity_cf'], column=col).value = 0

    ws.cell(row=exit_r['irr'], column=4).value = f"=IRR(D{exit_r['equity_cf']}:N{exit_r['equity_cf']})"
    ws.cell(row=exit_r['irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['irr'], column=4).number_format = PERCENT_FMT

    ws.cell(row=exit_r['moic'], column=4).value = \
        f"=(SUM(E{exit_r['equity_cf']}:N{exit_r['equity_cf']}))/-D{exit_r['equity_cf']}"
    ws.cell(row=exit_r['moic'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['moic'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['moic'], column=4).number_format = MULTIPLE_FMT

    # Unlevered
    ws.cell(row=exit_r['unlev_cf'], column=4).value = f"=-$D${inp['entry_ev']}-D{su['txn_fees']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = f"={c}{op['cfads']}"
        elif year == exit_year:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = f"={c}{op['cfads']}+D{exit_r['exit_ev']}"
        else:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = 0

    ws.cell(row=exit_r['unlev_irr'], column=4).value = f"=IRR(D{exit_r['unlev_cf']}:N{exit_r['unlev_cf']})"
    ws.cell(row=exit_r['unlev_irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['unlev_irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['unlev_irr'], column=4).number_format = PERCENT_FMT

    # ===== 8. AUDIT STRIP =====
    print("[8] Fixing Audit Strip...")

    ws.cell(row=3, column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=4, column=4).value = f"=E{op['ebitda']}"

    wb.save(filepath)
    print(f"\nSAVED: {filepath}")

    # Validation
    print("\n--- KEY FORMULAS ---")
    print(f"  D64 (Y1 Gen): {ws.cell(row=64, column=4).value}")
    print(f"  D65 (Debt): {ws.cell(row=65, column=4).value}")
    print(f"  D66 (ITC): {ws.cell(row=66, column=4).value}")
    print(f"  E69 (Gen Y1): {ws.cell(row=69, column=5).value}")
    print(f"  E78 (EBITDA): {ws.cell(row=78, column=5).value}")
    print(f"  D107 (Purchase): {ws.cell(row=107, column=4).value}")

if __name__ == "__main__":
    fix_solarbess()
