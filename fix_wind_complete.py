#!/usr/bin/env python3
"""
Complete fix for Model_6_Wind.xlsx
Key features: Generation degradation, PTC (Production Tax Credit), Sculpted debt
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'
CURRENCY_FMT = '#,##0.0_);(#,##0.0)'

def fix_wind():
    filepath = f"{DELIV_DIR}/Model_6_Wind.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    # ===== INPUT ROW MAPPING =====
    inp = {
        'exit_year_cp': 26,   # 7 (Control Panel)
        'ptc_end_year': 27,   # 5 (Control Panel)
        'entry_ev': 35,       # 180
        'txn_fees': 36,       # 0.015
        'exit_mult': 37,      # 8.5
        'capacity': 40,       # 100
        'cap_factor': 41,     # 0.38
        'degradation': 42,    # 0.003
        'ppa_price': 43,      # 45
        'merchant_price': 44, # 35
        'price_esc': 45,      # 0.02
        'om': 48,             # 35
        'om_esc': 49,         # 0.025
        'land_lease': 50,     # 0.5
        'insurance': 51,      # 0.003
        'tax_rate': 54,       # 0.25
        'ptc_rate': 55,       # 27.5
        'dep_basis': 56,      # 0.85
        'dep_life': 57,       # 20
        'leverage': 60,       # 0.65
        'int_rate': 61,       # 0.055
        'debt_tenor': 62,     # 7
        'target_dscr': 63,    # 1.35
        'dsra_months': 64,    # 6
    }

    calc = {
        'y1_gen': 68,  # Will place calculated items here
        'debt_amount': 69,
        'depreciation': 70,
    }

    op = {
        'generation': 73,
        'ppa_price': 74,
        'merchant_price': 75,
        'revenue': 77,
        'om': 79,
        'land_lease': 80,
        'insurance': 81,
        'ebitda': 83,
        'pre_ptc_tax': 86,
        'gross_ptc': 87,
        'ptc_applied': 88,
        'taxes': 89,
        'cfads': 91,
    }

    debt = {
        'beg_bal': 95,
        'interest': 96,
        'target_ds': 97,
        'sculpted_prin': 98,
        'actual_ds': 99,
        'end_bal': 100,
        'dscr': 101,
    }

    dsra = {
        'target': 105,
        'beg': 106,
        'funding': 107,
        'end': 108,
    }

    dist = {
        'avail': 111,
        'lockup': 112,
        'distributable': 113,
    }

    su = {
        'purchase': 117,
        'txn_fees': 118,
        'dsra_init': 119,
        'total_uses': 120,
        'debt': 123,
        'equity': 124,
        'total_sources': 125,
        'balance': 127,
    }

    exit_r = {
        'exit_ebitda': 130,
        'exit_ev': 131,
        'exit_debt': 132,
        'exit_dsra': 133,
        'exit_equity': 134,
        'equity_cf': 137,
        'irr': 138,
        'moic': 139,
        'unlev_cf': 142,
        'unlev_irr': 143,
    }

    print("="*60)
    print("FIXING MODEL_6_WIND.xlsx")
    print("="*60)

    # First, let's verify/find actual row numbers
    print("\n[0] Verifying structure...")
    for row in range(67, 145):
        a = ws.cell(row=row, column=1).value or ""
        if a and str(a).strip():
            print(f"  Row {row}: {str(a)[:40]}")

    # ===== 1. CALCULATED ITEMS =====
    print("\n[1] Fixing Calculated Items...")

    # Y1 Generation = Capacity * CF * 8760
    # Assume calc items are around row 68-70
    ws.cell(row=68, column=4).value = f"=$D${inp['capacity']}*$D${inp['cap_factor']}*8760"
    ws.cell(row=69, column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
    ws.cell(row=70, column=4).value = f"=$D${inp['entry_ev']}*$D${inp['dep_basis']}/$D${inp['dep_life']}"

    calc['y1_gen'] = 68
    calc['debt_amount'] = 69
    calc['depreciation'] = 70

    # ===== 2. OPERATING MODEL =====
    print("[2] Fixing Operating Model...")

    # Need to find actual operating model rows
    # Based on structure, let's assume standard layout
    op_rows = {
        'generation': 73,
        'ppa_price': 74,
        'merchant_price': 75,
        'revenue': 77,
        'om': 79,
        'land_lease': 80,
        'insurance': 81,
        'ebitda': 83,
        'pre_ptc_tax': 86,
        'gross_ptc': 87,
        'ptc_applied': 88,
        'taxes': 89,
        'cfads': 91,
    }

    ppa_end_year = 10  # Assume PPA covers full period (simplification)
    ptc_end_year = 5   # PTC ends Y5

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Generation with degradation
        ws.cell(row=op_rows['generation'], column=col).value = f"=$D$68*(1-$D${inp['degradation']})^({year}-1)"

        # PPA Price with escalation
        ws.cell(row=op_rows['ppa_price'], column=col).value = f"=$D${inp['ppa_price']}*(1+$D${inp['price_esc']})^({year}-1)"

        # Merchant Price with escalation
        ws.cell(row=op_rows['merchant_price'], column=col).value = f"=$D${inp['merchant_price']}*(1+$D${inp['price_esc']})^({year}-1)"

        # Revenue = Generation * Price / 1E6 (PPA price for simplicity)
        ws.cell(row=op_rows['revenue'], column=col).value = f"={c}{op_rows['generation']}*{c}{op_rows['ppa_price']}/1000000"

        # O&M with escalation ($/kW-yr)
        ws.cell(row=op_rows['om'], column=col).value = f"=$D${inp['capacity']}*$D${inp['om']}*(1+$D${inp['om_esc']})^({year}-1)/1000"

        # Land Lease
        ws.cell(row=op_rows['land_lease'], column=col).value = f"=$D${inp['land_lease']}*(1+$D${inp['price_esc']})^({year}-1)"

        # Insurance (% of EV)
        ws.cell(row=op_rows['insurance'], column=col).value = f"=$D${inp['entry_ev']}*$D${inp['insurance']}"

        # EBITDA = Revenue - O&M - Land - Insurance
        ws.cell(row=op_rows['ebitda'], column=col).value = \
            f"={c}{op_rows['revenue']}-{c}{op_rows['om']}-{c}{op_rows['land_lease']}-{c}{op_rows['insurance']}"
        ws.cell(row=op_rows['ebitda'], column=col).fill = GREEN_FILL
        ws.cell(row=op_rows['ebitda'], column=col).font = BLACK_BOLD
        ws.cell(row=op_rows['ebitda'], column=col).number_format = CURRENCY_FMT

        # Pre-PTC Tax = MAX(0, EBITDA - Depreciation) * Tax Rate
        ws.cell(row=op_rows['pre_ptc_tax'], column=col).value = \
            f"=MAX(0,{c}{op_rows['ebitda']}-$D$70)*$D${inp['tax_rate']}"

        # Gross PTC = IF(Year <= PTC End, Generation * PTC Rate / 1E6, 0)
        ws.cell(row=op_rows['gross_ptc'], column=col).value = \
            f"=IF({year}<=$D${inp['ptc_end_year']},{c}{op_rows['generation']}*$D${inp['ptc_rate']}/1000000,0)"

        # PTC Applied = MIN(Gross PTC, Pre-PTC Tax) - CANNOT create negative taxes
        ws.cell(row=op_rows['ptc_applied'], column=col).value = \
            f"=MIN({c}{op_rows['gross_ptc']},{c}{op_rows['pre_ptc_tax']})"

        # Taxes = Pre-PTC Tax - PTC Applied
        ws.cell(row=op_rows['taxes'], column=col).value = f"={c}{op_rows['pre_ptc_tax']}-{c}{op_rows['ptc_applied']}"

        # CFADS = EBITDA - Taxes
        ws.cell(row=op_rows['cfads'], column=col).value = f"={c}{op_rows['ebitda']}-{c}{op_rows['taxes']}"
        ws.cell(row=op_rows['cfads'], column=col).fill = GREEN_FILL
        ws.cell(row=op_rows['cfads'], column=col).font = BLACK_BOLD

    # ===== 3. SCULPTED DEBT SCHEDULE =====
    print("[3] Fixing Sculpted Debt Schedule...")

    debt_rows = {
        'beg_bal': 95,
        'interest': 96,
        'target_ds': 97,
        'sculpted_prin': 98,
        'actual_ds': 99,
        'end_bal': 100,
        'dscr': 101,
    }

    ws.cell(row=debt_rows['beg_bal'], column=4).value = 0
    ws.cell(row=debt_rows['end_bal'], column=4).value = f"=$D$69"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)

        if year == 1:
            ws.cell(row=debt_rows['beg_bal'], column=col).value = f"=$D$69"
        else:
            ws.cell(row=debt_rows['beg_bal'], column=col).value = f"={prev}{debt_rows['end_bal']}"

        ws.cell(row=debt_rows['interest'], column=col).value = f"={c}{debt_rows['beg_bal']}*$D${inp['int_rate']}"
        ws.cell(row=debt_rows['target_ds'], column=col).value = f"={c}{op_rows['cfads']}/$D${inp['target_dscr']}"
        ws.cell(row=debt_rows['sculpted_prin'], column=col).value = \
            f"=MIN(MAX(0,{c}{debt_rows['target_ds']}-{c}{debt_rows['interest']}),{c}{debt_rows['beg_bal']})"
        ws.cell(row=debt_rows['actual_ds'], column=col).value = f"={c}{debt_rows['interest']}+{c}{debt_rows['sculpted_prin']}"
        ws.cell(row=debt_rows['end_bal'], column=col).value = f"=MAX(0,{c}{debt_rows['beg_bal']}-{c}{debt_rows['sculpted_prin']})"
        ws.cell(row=debt_rows['dscr'], column=col).value = f"=IF({c}{debt_rows['actual_ds']}>0,{c}{op_rows['cfads']}/{c}{debt_rows['actual_ds']},0)"
        ws.cell(row=debt_rows['dscr'], column=col).number_format = MULTIPLE_FMT

    # ===== 4. DSRA =====
    print("[4] Fixing DSRA...")

    dsra_rows = {
        'target': 105,
        'beg': 106,
        'funding': 107,
        'end': 108,
    }

    ws.cell(row=dsra_rows['beg'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)
        next_c = get_column_letter(col + 1) if col < 14 else c

        if col < 14:
            ws.cell(row=dsra_rows['target'], column=col).value = f"={next_c}{debt_rows['actual_ds']}*$D${inp['dsra_months']}/12"
        else:
            ws.cell(row=dsra_rows['target'], column=col).value = 0

        if year == 1:
            ws.cell(row=dsra_rows['beg'], column=col).value = f"=D119"  # su['dsra_init']
        else:
            ws.cell(row=dsra_rows['beg'], column=col).value = f"={prev}{dsra_rows['end']}"

        ws.cell(row=dsra_rows['funding'], column=col).value = f"={c}{dsra_rows['target']}-{c}{dsra_rows['beg']}"
        ws.cell(row=dsra_rows['end'], column=col).value = f"={c}{dsra_rows['beg']}+{c}{dsra_rows['funding']}"

    # ===== 5. DISTRIBUTION WATERFALL =====
    print("[5] Fixing Distribution Waterfall...")

    dist_rows = {
        'avail': 111,
        'lockup': 112,
        'distributable': 113,
    }

    lockup_dscr = 1.1  # Hardcoded since not in inputs

    for col in range(5, 15):
        c = get_column_letter(col)

        ws.cell(row=dist_rows['avail'], column=col).value = \
            f"={c}{op_rows['cfads']}-{c}{debt_rows['actual_ds']}-{c}{dsra_rows['funding']}"
        ws.cell(row=dist_rows['lockup'], column=col).value = f"=IF({c}{debt_rows['dscr']}>=1.1,\"PASS\",\"LOCKED\")"
        ws.cell(row=dist_rows['distributable'], column=col).value = \
            f"=IF({c}{dist_rows['lockup']}=\"PASS\",MAX(0,{c}{dist_rows['avail']}),0)"

    # ===== 6. SOURCES & USES =====
    print("[6] Fixing Sources & Uses...")

    su_rows = {
        'purchase': 117,
        'txn_fees': 118,
        'dsra_init': 119,
        'total_uses': 120,
        'debt': 123,
        'equity': 124,
        'total_sources': 125,
        'balance': 127,
    }

    ws.cell(row=su_rows['purchase'], column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=su_rows['txn_fees'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['txn_fees']}"
    ws.cell(row=su_rows['dsra_init'], column=4).value = f"=E{dsra_rows['target']}"
    ws.cell(row=su_rows['total_uses'], column=4).value = f"=D{su_rows['purchase']}+D{su_rows['txn_fees']}+D{su_rows['dsra_init']}"

    ws.cell(row=su_rows['debt'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
    ws.cell(row=su_rows['equity'], column=4).value = f"=D{su_rows['total_uses']}-D{su_rows['debt']}"
    ws.cell(row=su_rows['total_sources'], column=4).value = f"=D{su_rows['debt']}+D{su_rows['equity']}"
    ws.cell(row=su_rows['balance'], column=4).value = f"=IF(ABS(D{su_rows['total_uses']}-D{su_rows['total_sources']})<0.01,\"PASS\",\"FAIL\")"

    # ===== 7. EXIT & RETURNS =====
    print("[7] Fixing Exit & Returns...")

    exit_year = 7
    exit_col = get_column_letter(4 + exit_year)

    exit_rows = {
        'exit_ebitda': 130,
        'exit_ev': 131,
        'exit_debt': 132,
        'exit_dsra': 133,
        'exit_equity': 134,
        'equity_cf': 137,
        'irr': 138,
        'moic': 139,
        'unlev_cf': 142,
        'unlev_irr': 143,
    }

    ws.cell(row=exit_rows['exit_ebitda'], column=4).value = f"={exit_col}{op_rows['ebitda']}"
    ws.cell(row=exit_rows['exit_ev'], column=4).value = f"=D{exit_rows['exit_ebitda']}*$D${inp['exit_mult']}"
    ws.cell(row=exit_rows['exit_debt'], column=4).value = f"={exit_col}{debt_rows['end_bal']}"
    ws.cell(row=exit_rows['exit_dsra'], column=4).value = f"={exit_col}{dsra_rows['end']}"
    ws.cell(row=exit_rows['exit_equity'], column=4).value = f"=D{exit_rows['exit_ev']}-D{exit_rows['exit_debt']}+D{exit_rows['exit_dsra']}"

    ws.cell(row=exit_rows['equity_cf'], column=4).value = f"=-D{su_rows['equity']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_rows['equity_cf'], column=col).value = f"={c}{dist_rows['distributable']}"
        elif year == exit_year:
            ws.cell(row=exit_rows['equity_cf'], column=col).value = f"={c}{dist_rows['distributable']}+D{exit_rows['exit_equity']}"
        else:
            ws.cell(row=exit_rows['equity_cf'], column=col).value = 0

    ws.cell(row=exit_rows['irr'], column=4).value = f"=IRR(D{exit_rows['equity_cf']}:N{exit_rows['equity_cf']})"
    ws.cell(row=exit_rows['irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_rows['irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_rows['irr'], column=4).number_format = PERCENT_FMT

    ws.cell(row=exit_rows['moic'], column=4).value = \
        f"=(SUM(E{exit_rows['equity_cf']}:K{exit_rows['equity_cf']}))/-D{exit_rows['equity_cf']}"
    ws.cell(row=exit_rows['moic'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_rows['moic'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_rows['moic'], column=4).number_format = MULTIPLE_FMT

    # Unlevered
    ws.cell(row=exit_rows['unlev_cf'], column=4).value = f"=-$D${inp['entry_ev']}-D{su_rows['txn_fees']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_rows['unlev_cf'], column=col).value = f"={c}{op_rows['cfads']}"
        elif year == exit_year:
            ws.cell(row=exit_rows['unlev_cf'], column=col).value = f"={c}{op_rows['cfads']}+D{exit_rows['exit_ev']}"
        else:
            ws.cell(row=exit_rows['unlev_cf'], column=col).value = 0

    ws.cell(row=exit_rows['unlev_irr'], column=4).value = f"=IRR(D{exit_rows['unlev_cf']}:N{exit_rows['unlev_cf']})"
    ws.cell(row=exit_rows['unlev_irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_rows['unlev_irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_rows['unlev_irr'], column=4).number_format = PERCENT_FMT

    # ===== 8. AUDIT STRIP =====
    print("[8] Fixing Audit Strip...")

    ws.cell(row=3, column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=4, column=4).value = f"=E{op_rows['ebitda']}"

    wb.save(filepath)
    print(f"\nSAVED: {filepath}")

    print("\n--- KEY FORMULAS ---")
    print(f"  D68 (Y1 Gen): {ws.cell(row=68, column=4).value}")
    print(f"  E73 (Gen Y1): {ws.cell(row=73, column=5).value}")
    print(f"  E77 (Revenue): {ws.cell(row=77, column=5).value}")
    print(f"  E83 (EBITDA): {ws.cell(row=83, column=5).value}")
    print(f"  E88 (PTC Applied): {ws.cell(row=88, column=5).value}")

if __name__ == "__main__":
    fix_wind()
