#!/usr/bin/env python3
"""
Build Model 2: Peaker (Simple Cycle Gas Turbine)
Infrastructure PE Interview Prep - Lotus Infrastructure Partners

Key differences from CCGT:
- Uses dispatch hours (400) instead of capacity factor
- Lower leverage (30%)
- Lower cash sweep (50%)
- Major maintenance expense in Y5 ($4mm)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# FORMATTING CONSTANTS (same as CCGT)
# ============================================================================
YELLOW_FILL = PatternFill(start_color="FFFFC7", end_color="FFFFC7", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
NAVY_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

BLUE_BOLD = Font(color="0000FF", bold=True)
BLACK_FONT = Font(color="000000")
BLACK_BOLD = Font(color="000000", bold=True)
WHITE_BOLD = Font(color="FFFFFF", bold=True)

THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

CURRENCY_FORMAT = '#,##0.0_);(#,##0.0)'
PERCENT_FORMAT = '0.0%'
MULTIPLE_FORMAT = '0.00"x"'
INTEGER_FORMAT = '#,##0'

def col_letter(col_num):
    return get_column_letter(col_num)

def apply_input_style(ws, row, col, value=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    cell.fill = YELLOW_FILL
    cell.font = BLUE_BOLD
    cell.border = THIN_BORDER
    return cell

def apply_calc_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = PatternFill(fill_type=None)
    cell.font = BLACK_FONT
    cell.border = THIN_BORDER
    return cell

def apply_output_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = GREEN_FILL
    cell.font = BLACK_BOLD
    cell.border = THIN_BORDER
    return cell

def apply_header_style(ws, row, col, text):
    cell = ws.cell(row=row, column=col)
    cell.value = text
    cell.fill = NAVY_FILL
    cell.font = WHITE_BOLD
    cell.alignment = Alignment(horizontal='left')
    return cell

def apply_warning_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    cell.fill = ORANGE_FILL
    cell.font = BLACK_FONT
    cell.border = THIN_BORDER
    return cell

def set_label(ws, row, label, units="", notes=""):
    ws.cell(row=row, column=1, value=label).font = BLACK_FONT
    ws.cell(row=row, column=2, value=units).font = Font(color="666666", italic=True)
    ws.cell(row=row, column=3, value=notes).font = Font(color="666666", italic=True)

def build_peaker_model(output_path):
    """Build the Peaker model"""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"

    rows = {}

    # Column widths
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 38
    for col in range(4, 15):
        ws.column_dimensions[col_letter(col)].width = 14

    # ========================================================================
    # ROW 1: YEAR HEADERS
    # ========================================================================
    row = 1
    rows['year_header'] = row
    ws.cell(row=row, column=1, value="Peaker Model - Lotus Infrastructure Partners").font = Font(bold=True, size=14)
    ws.cell(row=row, column=4, value=0)
    for yr in range(1, 11):
        ws.cell(row=row, column=4+yr, value=yr)
    for col in range(4, 15):
        cell = ws.cell(row=row, column=col)
        cell.font = BLACK_BOLD
        cell.alignment = Alignment(horizontal='center')

    # ========================================================================
    # AUDIT STRIP (Rows 2-13)
    # ========================================================================
    row = 2
    rows['audit_header'] = row
    apply_header_style(ws, row, 1, "═══ AUDIT STRIP ═══")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 3; rows['audit_entry_ev'] = row
    set_label(ws, row, "Entry EV", "$mm")

    row = 4; rows['audit_y1_ebitda'] = row
    set_label(ws, row, "Y1 EBITDA", "$mm")

    row = 5; rows['audit_avg_cfads'] = row
    set_label(ws, row, "Avg CFADS (Y1-Y7)", "$mm")

    row = 6; rows['audit_min_dscr'] = row
    set_label(ws, row, "Min DSCR (Y1-Y7)", "x")

    row = 7; rows['audit_levered_irr'] = row
    set_label(ws, row, "Levered IRR", "%")

    row = 8; rows['audit_moic'] = row
    set_label(ws, row, "MOIC", "x")

    row = 9; rows['audit_unlevered_irr'] = row
    set_label(ws, row, "Unlevered IRR", "%")

    row = 10; rows['qa_header'] = row
    apply_header_style(ws, row, 1, "═══ QA CHECKS ═══")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 11; rows['qa_su_balance'] = row
    set_label(ws, row, "S&U Balance")

    row = 12; rows['qa_debt_payoff'] = row
    set_label(ws, row, "Debt Payoff Y7")

    row = 13; rows['qa_min_dscr_check'] = row
    set_label(ws, row, "Min DSCR ≥ 1.25x")

    row = 15

    # ========================================================================
    # INPUTS SECTION
    # ========================================================================

    # Transaction Inputs
    row = 16; rows['input_header_txn'] = row
    apply_header_style(ws, row, 1, "TRANSACTION INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 17; rows['entry_ev'] = row
    set_label(ws, row, "Entry EV", "$mm", "Purchase price")
    apply_input_style(ws, row, 4, 110)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 18; rows['txn_fees'] = row
    set_label(ws, row, "Transaction Fees", "%", "Legal, advisory, financing")
    apply_input_style(ws, row, 4, 0.015)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 19; rows['exit_multiple'] = row
    set_label(ws, row, "Exit Multiple", "x", "On exit year EBITDA")
    apply_input_style(ws, row, 4, 7.0)
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    row = 20; rows['exit_year'] = row
    set_label(ws, row, "Exit Year", "#", "Hold period")
    apply_input_style(ws, row, 4, 7)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 22
    # Asset Inputs
    rows['input_header_asset'] = row
    apply_header_style(ws, row, 1, "ASSET INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 23; rows['capacity'] = row
    set_label(ws, row, "Capacity", "MW", "Nameplate rating")
    apply_input_style(ws, row, 4, 200)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 24; rows['heat_rate'] = row
    set_label(ws, row, "Heat Rate", "Btu/kWh", "Less efficient than CCGT")
    apply_input_style(ws, row, 4, 10500)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 25; rows['availability'] = row
    set_label(ws, row, "Availability", "%", "For UCAP accreditation")
    apply_input_style(ws, row, 4, 0.94)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 26; rows['forced_outage_rate'] = row
    set_label(ws, row, "Forced Outage Rate", "%", "Unplanned downtime")
    apply_input_style(ws, row, 4, 0.05)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 27; rows['dispatch_hours'] = row
    set_label(ws, row, "Dispatch Hours", "hours/yr", "Peak hours only (~400)")
    apply_input_style(ws, row, 4, 400)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 29
    # Pricing Inputs
    rows['input_header_pricing'] = row
    apply_header_style(ws, row, 1, "PRICING INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 30; rows['power_price_y1'] = row
    set_label(ws, row, "Power Price Y1", "$/MWh", "Higher (dispatch-weighted)")
    apply_input_style(ws, row, 4, 95)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 31; rows['capacity_price_y1'] = row
    set_label(ws, row, "Capacity Price Y1", "$/kW-yr", "ICAP/RPM result")
    apply_input_style(ws, row, 4, 60)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 32; rows['gas_price_y1'] = row
    set_label(ws, row, "Gas Price Y1", "$/mmBtu", "Henry Hub + basis")
    apply_input_style(ws, row, 4, 3.25)
    ws.cell(row=row, column=4).number_format = '#,##0.00'

    row = 33; rows['escalation'] = row
    set_label(ws, row, "Escalation", "%/yr", "Applied to all prices/costs")
    apply_input_style(ws, row, 4, 0.025)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 35
    # Cost Inputs
    rows['input_header_costs'] = row
    apply_header_style(ws, row, 1, "COST INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 36; rows['vom'] = row
    set_label(ws, row, "VOM", "$/MWh", "Variable O&M")
    apply_input_style(ws, row, 4, 4.00)
    ws.cell(row=row, column=4).number_format = '#,##0.00'

    row = 37; rows['fom'] = row
    set_label(ws, row, "FOM", "$/kW-yr", "Fixed O&M")
    apply_input_style(ws, row, 4, 12.00)
    ws.cell(row=row, column=4).number_format = '#,##0.00'

    row = 38; rows['maint_year'] = row
    set_label(ws, row, "Major Maintenance Year", "#", "Hot gas path inspection")
    apply_input_style(ws, row, 4, 5)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 39; rows['maint_cost'] = row
    set_label(ws, row, "Major Maintenance Cost", "$mm", "HGP inspection cost")
    apply_input_style(ws, row, 4, 4)
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 41
    # Financing Inputs
    rows['input_header_fin'] = row
    apply_header_style(ws, row, 1, "FINANCING INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 42; rows['leverage'] = row
    set_label(ws, row, "Leverage", "%", "Debt / Entry EV (lower)")
    apply_input_style(ws, row, 4, 0.30)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 43; rows['interest_rate'] = row
    set_label(ws, row, "Interest Rate", "%", "All-in borrowing cost")
    apply_input_style(ws, row, 4, 0.07)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 44; rows['debt_tenor'] = row
    set_label(ws, row, "Debt Tenor", "years", "Amortization period")
    apply_input_style(ws, row, 4, 7)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 45; rows['dsra_months'] = row
    set_label(ws, row, "DSRA Months", "months", "Debt service reserve")
    apply_input_style(ws, row, 4, 6)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 46; rows['sweep_pct'] = row
    set_label(ws, row, "Cash Sweep %", "%", "Lower for maint reserves")
    apply_input_style(ws, row, 4, 0.50)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 47; rows['lockup_dscr'] = row
    set_label(ws, row, "Lock-up DSCR", "x", "Distribution threshold")
    apply_input_style(ws, row, 4, 1.10)
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    row = 49
    # Tax Inputs
    rows['input_header_tax'] = row
    apply_header_style(ws, row, 1, "TAX INPUTS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 50; rows['tax_rate'] = row
    set_label(ws, row, "Tax Rate", "%", "Blended federal + state")
    apply_input_style(ws, row, 4, 0.25)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 51; rows['depreciable_basis'] = row
    set_label(ws, row, "Depreciable Basis", "%", "% of EV for depreciation")
    apply_input_style(ws, row, 4, 0.85)
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 52; rows['depreciation_life'] = row
    set_label(ws, row, "Depreciation Life", "years", "Straight-line")
    apply_input_style(ws, row, 4, 15)
    ws.cell(row=row, column=4).number_format = INTEGER_FORMAT

    row = 54
    # ========================================================================
    # CALCULATED ITEMS
    # ========================================================================
    rows['calc_header'] = row
    apply_header_style(ws, row, 1, "CALCULATED ITEMS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 55; rows['ucap'] = row
    set_label(ws, row, "UCAP (Unforced Capacity)", "MW", "Capacity × (1 - FOR)")
    apply_calc_style(ws, row, 4, f"=$D${rows['capacity']}*(1-$D${rows['forced_outage_rate']})")
    ws.cell(row=row, column=4).number_format = '#,##0.0'

    row = 56; rows['generation'] = row
    set_label(ws, row, "Annual Generation", "MWh", "Capacity × Hours × Avail")
    # PEAKER: Uses dispatch hours, NOT capacity factor × 8760
    apply_calc_style(ws, row, 4, f"=$D${rows['capacity']}*$D${rows['dispatch_hours']}*$D${rows['availability']}")
    ws.cell(row=row, column=4).number_format = '#,##0'

    row = 57; rows['debt_amount'] = row
    set_label(ws, row, "Debt Amount", "$mm", "Entry EV × Leverage")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['leverage']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 58; rows['annual_depreciation'] = row
    set_label(ws, row, "Annual Depreciation", "$mm", "EV × Basis / Life")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['depreciable_basis']}/$D${rows['depreciation_life']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 59; rows['scheduled_principal'] = row
    set_label(ws, row, "Scheduled Principal", "$mm/yr", "Debt / Tenor")
    apply_calc_style(ws, row, 4, f"=$D${rows['debt_amount']}/$D${rows['debt_tenor']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 61
    # ========================================================================
    # OPERATING MODEL
    # ========================================================================
    rows['op_header'] = row
    apply_header_style(ws, row, 1, "OPERATING MODEL")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 62; rows['power_price'] = row
    set_label(ws, row, "Power Price", "$/MWh", "Escalated")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['power_price_y1']}*(1+$D${rows['escalation']})^({yr}-1)"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 63; rows['capacity_price'] = row
    set_label(ws, row, "Capacity Price", "$/kW-yr", "Escalated")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['capacity_price_y1']}*(1+$D${rows['escalation']})^({yr}-1)"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 64; rows['gas_price'] = row
    set_label(ws, row, "Gas Price", "$/mmBtu", "Escalated")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['gas_price_y1']}*(1+$D${rows['escalation']})^({yr}-1)"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = '#,##0.00'

    row = 66; rows['energy_revenue'] = row
    set_label(ws, row, "Energy Revenue", "$mm", "Generation × Power Price / 1M")
    for col in range(5, 15):
        formula = f"=$D${rows['generation']}*{col_letter(col)}{rows['power_price']}/1000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 67; rows['capacity_revenue'] = row
    set_label(ws, row, "Capacity Revenue", "$mm", "UCAP × Cap Price × 1000 / 1M")
    for col in range(5, 15):
        formula = f"=$D${rows['ucap']}*{col_letter(col)}{rows['capacity_price']}*1000/1000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 68; rows['total_revenue'] = row
    set_label(ws, row, "Total Revenue", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['energy_revenue']}+{col_letter(col)}{rows['capacity_revenue']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
        ws.cell(row=row, column=col).font = BLACK_BOLD

    row = 70; rows['fuel_cost'] = row
    set_label(ws, row, "Fuel Cost", "$mm", "Gen × HR × Gas / 1B")
    for col in range(5, 15):
        formula = f"=$D${rows['generation']}*$D${rows['heat_rate']}*{col_letter(col)}{rows['gas_price']}/1000000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 71; rows['vom_cost'] = row
    set_label(ws, row, "VOM", "$mm", "Generation × VOM rate / 1M")
    for col in range(5, 15):
        formula = f"=$D${rows['generation']}*$D${rows['vom']}/1000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 72; rows['fom_cost'] = row
    set_label(ws, row, "FOM", "$mm", "Capacity × FOM × esc × 1000 / 1M")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=$D${rows['capacity']}*$D${rows['fom']}*(1+$D${rows['escalation']})^({yr}-1)*1000/1000000"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 73; rows['total_opex'] = row
    set_label(ws, row, "Total Operating Costs", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['fuel_cost']}+{col_letter(col)}{rows['vom_cost']}+{col_letter(col)}{rows['fom_cost']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 75; rows['ebitda'] = row
    set_label(ws, row, "EBITDA", "$mm")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['total_revenue']}-{col_letter(col)}{rows['total_opex']}"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 77
    # ========================================================================
    # CASH FLOW
    # ========================================================================
    rows['cf_header'] = row
    apply_header_style(ws, row, 1, "CASH FLOW")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 78; rows['ebit'] = row
    set_label(ws, row, "EBIT", "$mm", "EBITDA - Depreciation")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['ebitda']}-$D${rows['annual_depreciation']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 79; rows['taxes'] = row
    set_label(ws, row, "Taxes", "$mm", "MAX(0, EBIT) × Tax Rate")
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['ebit']})*$D${rows['tax_rate']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 80; rows['major_maint'] = row
    set_label(ws, row, "Major Maintenance", "$mm", "HGP in specified year")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=IF({yr}=$D${rows['maint_year']},$D${rows['maint_cost']},0)"
        apply_warning_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 81; rows['cfads'] = row
    set_label(ws, row, "CFADS", "$mm", "EBITDA - Taxes - Maint")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['ebitda']}-{col_letter(col)}{rows['taxes']}-{col_letter(col)}{rows['major_maint']}"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 83
    # ========================================================================
    # DEBT SCHEDULE (50% Cash Sweep)
    # ========================================================================
    rows['debt_header'] = row
    apply_header_style(ws, row, 1, "DEBT SCHEDULE (50% Cash Sweep)")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 84; rows['debt_beg_bal'] = row
    set_label(ws, row, "Beginning Balance", "$mm")
    apply_calc_style(ws, row, 5, f"=$D${rows['debt_amount']}")
    ws.cell(row=row, column=5).number_format = CURRENCY_FORMAT
    for col in range(6, 15):
        prev_col = col_letter(col - 1)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 85; rows['interest'] = row
    set_label(ws, row, "Interest", "$mm", "Beg Bal × Int Rate")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['debt_beg_bal']}*$D${rows['interest_rate']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 86; rows['sched_principal'] = row
    set_label(ws, row, "Scheduled Principal", "$mm", "MIN(Sched, Beg Bal)")
    for col in range(5, 15):
        formula = f"=MIN($D${rows['scheduled_principal']},{col_letter(col)}{rows['debt_beg_bal']})"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 87; rows['debt_service_pre'] = row
    set_label(ws, row, "Debt Service (Pre-Sweep)", "$mm", "Interest + Sched Prin")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['interest']}+{col_letter(col)}{rows['sched_principal']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 88; rows['excess_cash'] = row
    set_label(ws, row, "Excess Cash", "$mm", "CFADS - DS - DSRA Funding")

    row = 89; rows['cash_sweep'] = row
    set_label(ws, row, "Cash Sweep", "$mm", "MIN(Excess×50%, Beg-Sched)")
    for col in range(5, 15):
        formula = f"=MIN(MAX(0,{col_letter(col)}{rows['excess_cash']})*$D${rows['sweep_pct']},MAX(0,{col_letter(col)}{rows['debt_beg_bal']}-{col_letter(col)}{rows['sched_principal']}))"
        apply_warning_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 90; rows['total_principal'] = row
    set_label(ws, row, "Total Principal", "$mm", "Scheduled + Sweep")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['sched_principal']}+{col_letter(col)}{rows['cash_sweep']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 91; rows['debt_end_bal'] = row
    set_label(ws, row, "Ending Balance", "$mm", "MAX(0, Beg - Total Prin)")
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['debt_beg_bal']}-{col_letter(col)}{rows['total_principal']})"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    # Fix beginning balance references
    for col in range(6, 15):
        prev_col = col_letter(col - 1)
        formula = f"={prev_col}{rows['debt_end_bal']}"
        apply_calc_style(ws, rows['debt_beg_bal'], col, formula)

    row = 92; rows['dscr'] = row
    set_label(ws, row, "DSCR", "x", "CFADS / Debt Service")
    for col in range(5, 15):
        formula = f"=IF({col_letter(col)}{rows['debt_service_pre']}>0,{col_letter(col)}{rows['cfads']}/{col_letter(col)}{rows['debt_service_pre']},0)"
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = MULTIPLE_FORMAT

    # Add Scheduled DS row (for DSRA - breaks circularity)
    row = 93; rows['scheduled_ds'] = row
    set_label(ws, row, "Scheduled DS (for DSRA)", "$mm", "Based on straight-line amort")
    for col in range(5, 15):
        yr = col - 4
        formula = f"=MAX(0,$D${rows['debt_amount']}-({yr}-1)*$D${rows['scheduled_principal']})*$D${rows['interest_rate']}+$D${rows['scheduled_principal']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 94
    # ========================================================================
    # DSRA
    # ========================================================================
    rows['dsra_header'] = row
    apply_header_style(ws, row, 1, "DSRA (Debt Service Reserve Account)")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 95; rows['dsra_target'] = row
    set_label(ws, row, "DSRA Target", "$mm", "Next Yr Scheduled DS × DSRA Months / 12")
    apply_calc_style(ws, row, 4, f"=E{rows['scheduled_ds']}*$D${rows['dsra_months']}/12")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 14):
        next_col = col_letter(col + 1)
        formula = f"={next_col}{rows['scheduled_ds']}*$D${rows['dsra_months']}/12"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
    apply_calc_style(ws, row, 14, "=0")
    ws.cell(row=row, column=14).number_format = CURRENCY_FORMAT

    row = 96; rows['dsra_beg'] = row
    set_label(ws, row, "DSRA Beginning", "$mm")
    apply_calc_style(ws, row, 4, "=0")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 15):
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 97; rows['dsra_funding'] = row
    set_label(ws, row, "DSRA Funding/(Release)", "$mm", "Target - Beginning")
    for col in range(4, 15):
        formula = f"={col_letter(col)}{rows['dsra_target']}-{col_letter(col)}{rows['dsra_beg']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 98; rows['dsra_end'] = row
    set_label(ws, row, "DSRA Ending", "$mm", "Beginning + Funding")
    for col in range(4, 15):
        formula = f"={col_letter(col)}{rows['dsra_beg']}+{col_letter(col)}{rows['dsra_funding']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    # Fix DSRA Beginning references
    for col in range(5, 15):
        prev_col = col_letter(col - 1)
        formula = f"={prev_col}{rows['dsra_end']}"
        apply_calc_style(ws, rows['dsra_beg'], col, formula)

    # Add Excess Cash formulas
    for col in range(5, 15):
        formula = f"=MAX(0,{col_letter(col)}{rows['cfads']}-{col_letter(col)}{rows['debt_service_pre']}-{col_letter(col)}{rows['dsra_funding']})"
        apply_calc_style(ws, rows['excess_cash'], col, formula)
        ws.cell(row=rows['excess_cash'], column=col).number_format = CURRENCY_FORMAT

    row = 100
    # ========================================================================
    # DISTRIBUTION WATERFALL
    # ========================================================================
    rows['dist_header'] = row
    apply_header_style(ws, row, 1, "DISTRIBUTION WATERFALL")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 101; rows['cash_available'] = row
    set_label(ws, row, "Cash Available for Distribution", "$mm", "CFADS - Total DS - DSRA Fund")
    for col in range(5, 15):
        formula = f"={col_letter(col)}{rows['cfads']}-({col_letter(col)}{rows['interest']}+{col_letter(col)}{rows['total_principal']})-{col_letter(col)}{rows['dsra_funding']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 102; rows['lockup_test'] = row
    set_label(ws, row, "Lock-up Test", "", "PASS if DSCR ≥ Lock-up")
    for col in range(5, 15):
        formula = f'=IF({col_letter(col)}{rows["dscr"]}>=$D${rows["lockup_dscr"]},"PASS","LOCKED")'
        apply_calc_style(ws, row, col, formula)

    row = 103; rows['distributable'] = row
    set_label(ws, row, "Distributable Cash", "$mm", "If PASS, MAX(0, Cash Avail)")
    for col in range(5, 15):
        formula = f'=IF({col_letter(col)}{rows["lockup_test"]}="PASS",MAX(0,{col_letter(col)}{rows["cash_available"]}),0)'
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 105
    # ========================================================================
    # SOURCES & USES
    # ========================================================================
    rows['su_header'] = row
    apply_header_style(ws, row, 1, "SOURCES & USES (Year 0)")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 106; rows['uses_header'] = row
    set_label(ws, row, "USES", "", "")
    ws.cell(row=row, column=1).font = BLACK_BOLD

    row = 107; rows['use_purchase'] = row
    set_label(ws, row, "  Purchase Price", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 108; rows['use_fees'] = row
    set_label(ws, row, "  Transaction Fees", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['entry_ev']}*$D${rows['txn_fees']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 109; rows['use_dsra'] = row
    set_label(ws, row, "  DSRA Funding", "$mm")
    apply_calc_style(ws, row, 4, f"=D{rows['dsra_target']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 110; rows['total_uses'] = row
    set_label(ws, row, "Total Uses", "$mm")
    apply_output_style(ws, row, 4, f"=D{rows['use_purchase']}+D{rows['use_fees']}+D{rows['use_dsra']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 112; rows['sources_header'] = row
    set_label(ws, row, "SOURCES", "", "")
    ws.cell(row=row, column=1).font = BLACK_BOLD

    row = 113; rows['source_debt'] = row
    set_label(ws, row, "  Debt", "$mm")
    apply_calc_style(ws, row, 4, f"=$D${rows['debt_amount']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 114; rows['source_equity'] = row
    set_label(ws, row, "  Equity", "$mm", "Plug to balance")
    apply_calc_style(ws, row, 4, f"=D{rows['total_uses']}-D{rows['source_debt']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 115; rows['total_sources'] = row
    set_label(ws, row, "Total Sources", "$mm")
    apply_output_style(ws, row, 4, f"=D{rows['source_debt']}+D{rows['source_equity']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT

    row = 116; rows['su_check'] = row
    set_label(ws, row, "Balance Check", "", "")
    apply_output_style(ws, row, 4, f'=IF(ABS(D{rows["total_uses"]}-D{rows["total_sources"]})<0.01,"PASS","FAIL")')

    row = 118
    # ========================================================================
    # EXIT & EQUITY RETURNS
    # ========================================================================
    rows['exit_header'] = row
    apply_header_style(ws, row, 1, "EXIT CALCULATIONS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 119; rows['exit_ev'] = row
    set_label(ws, row, "Exit EV", "$mm", "Exit Multiple × Exit Yr EBITDA")
    apply_calc_style(ws, row, 11, f"=$D${rows['exit_multiple']}*K{rows['ebitda']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 120; rows['exit_debt'] = row
    set_label(ws, row, "Exit Debt Payoff", "$mm", "Ending Bal at Exit")
    apply_calc_style(ws, row, 11, f"=K{rows['debt_end_bal']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 121; rows['exit_dsra_release'] = row
    set_label(ws, row, "DSRA Release", "$mm", "DSRA Ending at Exit")
    apply_calc_style(ws, row, 11, f"=K{rows['dsra_end']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 122; rows['exit_equity_proceeds'] = row
    set_label(ws, row, "Exit Equity Proceeds", "$mm", "Exit EV - Debt + DSRA")
    apply_output_style(ws, row, 11, f"=K{rows['exit_ev']}-K{rows['exit_debt']}+K{rows['exit_dsra_release']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT

    row = 124; rows['returns_header'] = row
    apply_header_style(ws, row, 1, "EQUITY RETURNS")
    for col in range(2, 15):
        ws.cell(row=row, column=col).fill = NAVY_FILL

    row = 125; rows['equity_cf'] = row
    set_label(ws, row, "Equity Cash Flow", "$mm")
    apply_calc_style(ws, row, 4, f"=-D{rows['source_equity']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 11):
        formula = f"={col_letter(col)}{rows['distributable']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
    apply_calc_style(ws, row, 11, f"=K{rows['distributable']}+K{rows['exit_equity_proceeds']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT
    for col in range(12, 15):
        apply_calc_style(ws, row, col, "=0")
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 126; rows['levered_irr'] = row
    set_label(ws, row, "Levered IRR", "%")
    apply_output_style(ws, row, 4, f"=IRR(D{rows['equity_cf']}:K{rows['equity_cf']})")
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    row = 127; rows['moic'] = row
    set_label(ws, row, "MOIC", "x", "Sum inflows / outflow")
    apply_output_style(ws, row, 4, f"=SUMIF(D{rows['equity_cf']}:K{rows['equity_cf']},\">0\")/ABS(D{rows['equity_cf']})")
    ws.cell(row=row, column=4).number_format = MULTIPLE_FORMAT

    row = 129; rows['unlev_header'] = row
    set_label(ws, row, "Unlevered Returns (for reference)", "")
    ws.cell(row=row, column=1).font = BLACK_BOLD

    row = 130; rows['unlev_cf'] = row
    set_label(ws, row, "Unlevered Cash Flow", "$mm")
    apply_calc_style(ws, row, 4, f"=-$D${rows['entry_ev']}-D{rows['use_fees']}")
    ws.cell(row=row, column=4).number_format = CURRENCY_FORMAT
    for col in range(5, 11):
        formula = f"={col_letter(col)}{rows['cfads']}"
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT
    apply_calc_style(ws, row, 11, f"=K{rows['cfads']}+K{rows['exit_ev']}")
    ws.cell(row=row, column=11).number_format = CURRENCY_FORMAT
    for col in range(12, 15):
        apply_calc_style(ws, row, col, "=0")
        ws.cell(row=row, column=col).number_format = CURRENCY_FORMAT

    row = 131; rows['unlev_irr'] = row
    set_label(ws, row, "Unlevered IRR", "%")
    apply_output_style(ws, row, 4, f"=IRR(D{rows['unlev_cf']}:K{rows['unlev_cf']})")
    ws.cell(row=row, column=4).number_format = PERCENT_FORMAT

    # ========================================================================
    # POPULATE AUDIT STRIP
    # ========================================================================
    apply_output_style(ws, rows['audit_entry_ev'], 4, f"=$D${rows['entry_ev']}")
    ws.cell(row=rows['audit_entry_ev'], column=4).number_format = CURRENCY_FORMAT

    apply_output_style(ws, rows['audit_y1_ebitda'], 4, f"=E{rows['ebitda']}")
    ws.cell(row=rows['audit_y1_ebitda'], column=4).number_format = CURRENCY_FORMAT

    apply_output_style(ws, rows['audit_avg_cfads'], 4, f"=AVERAGE(E{rows['cfads']}:K{rows['cfads']})")
    ws.cell(row=rows['audit_avg_cfads'], column=4).number_format = CURRENCY_FORMAT

    apply_output_style(ws, rows['audit_min_dscr'], 4, f"=MIN(E{rows['dscr']}:K{rows['dscr']})")
    ws.cell(row=rows['audit_min_dscr'], column=4).number_format = MULTIPLE_FORMAT

    apply_output_style(ws, rows['audit_levered_irr'], 4, f"=D{rows['levered_irr']}")
    ws.cell(row=rows['audit_levered_irr'], column=4).number_format = PERCENT_FORMAT

    apply_output_style(ws, rows['audit_moic'], 4, f"=D{rows['moic']}")
    ws.cell(row=rows['audit_moic'], column=4).number_format = MULTIPLE_FORMAT

    apply_output_style(ws, rows['audit_unlevered_irr'], 4, f"=D{rows['unlev_irr']}")
    ws.cell(row=rows['audit_unlevered_irr'], column=4).number_format = PERCENT_FORMAT

    apply_output_style(ws, rows['qa_su_balance'], 4, f"=D{rows['su_check']}")
    apply_output_style(ws, rows['qa_debt_payoff'], 4, f'=IF(K{rows["debt_end_bal"]}<1,"PASS","FAIL")')
    apply_output_style(ws, rows['qa_min_dscr_check'], 4, f'=IF(D{rows["audit_min_dscr"]}>=1.25,"PASS","FAIL")')

    # Freeze panes
    ws.freeze_panes = 'E2'

    wb.save(output_path)
    print(f"Peaker model saved to: {output_path}")
    return rows

if __name__ == "__main__":
    output_path = "/mnt/user-data/outputs/Model_2_Peaker.xlsx"
    rows = build_peaker_model(output_path)
    print("Peaker model built successfully!")
