#!/usr/bin/env python3
"""
Corrected fix for Model_6_Wind.xlsx based on actual row structure
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'
CURRENCY_FMT = '#,##0.0_);(#,##0.0)'

def fix_wind():
    filepath = f"{DELIV_DIR}/Model_6_Wind.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    # ===== ACTUAL ROW MAPPING (verified from printout) =====
    inp = {
        'exit_year_cp': 26,
        'ptc_end_year': 27,
        'entry_ev': 35,
        'txn_fees': 36,
        'exit_mult': 37,
        'capacity': 40,
        'cap_factor': 41,
        'degradation': 42,
        'ppa_price': 43,
        'merchant_price': 44,
        'price_esc': 45,
        'om': 48,
        'om_esc': 49,
        'land_lease': 50,
        'insurance': 51,
        'tax_rate': 54,
        'ptc_rate': 55,
        'dep_basis': 56,
        'dep_life': 57,
        'leverage': 60,
        'int_rate': 61,
        'debt_tenor': 62,
        'target_dscr': 63,
        'dsra_months': 64,
    }

    # ACTUAL operating model rows
    op = {
        'generation': 68,
        'revenue': 69,
        'om': 71,
        'land_lease': 72,
        'insurance': 73,
        'ebitda': 74,
        'depreciation': 77,
        'ebit': 78,
        'pre_ptc_tax': 79,
        'gross_ptc': 80,
        'ptc_applied': 81,
        'taxes': 82,
        'cfads': 84,
    }

    # ACTUAL debt schedule rows
    debt = {
        'target_ds': 87,
        'beg_bal': 89,
        'interest': 90,
        'sculpted_prin': 91,
        'capped_prin': 92,
        'end_bal': 93,
        'actual_ds': 95,
        'dscr': 96,
    }

    dsra = {
        'target': 100,
        'beg': 101,
        'funding': 102,
        'end': 103,
    }

    dist = {
        'avail': 106,
        'distributable': 107,
    }

    su = {
        'purchase': 111,
        'txn_fees': 112,
        'dsra_init': 113,
        'total_uses': 114,
        'debt': 117,
        'equity': 118,
        'total_sources': 119,
        'balance': 120,
    }

    exit_r = {
        'exit_ebitda': 123,
        'exit_ev': 124,
        'exit_debt': 125,
        'exit_dsra': 126,
        'exit_equity': 127,
        'equity_cf': 129,
        'irr': 131,
        'moic': 132,
    }

    print("="*60)
    print("FIXING MODEL_6_WIND.xlsx (CORRECTED)")
    print("="*60)

    # ===== 1. CALCULATED ITEMS =====
    print("[1] Fixing Calculated Items (inline)...")

    # Y1 Generation stored in a helper cell or calculated inline
    # Debt Amount and Depreciation

    # ===== 2. OPERATING MODEL =====
    print("[2] Fixing Operating Model...")

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Generation with degradation: Capacity * CF * 8760 * (1-deg)^(yr-1)
        ws.cell(row=op['generation'], column=col).value = \
            f"=$D${inp['capacity']}*$D${inp['cap_factor']}*8760*(1-$D${inp['degradation']})^({year}-1)"

        # Revenue = Generation * PPA Price / 1E6
        # PPA Price with escalation
        ws.cell(row=op['revenue'], column=col).value = \
            f"={c}{op['generation']}*$D${inp['ppa_price']}*(1+$D${inp['price_esc']})^({year}-1)/1000000"

        # O&M ($/kW-yr * capacity / 1000 to get $mm)
        ws.cell(row=op['om'], column=col).value = \
            f"=$D${inp['capacity']}*$D${inp['om']}*(1+$D${inp['om_esc']})^({year}-1)/1000"

        # Land Lease
        ws.cell(row=op['land_lease'], column=col).value = \
            f"=$D${inp['land_lease']}*(1+$D${inp['price_esc']})^({year}-1)"

        # Insurance
        ws.cell(row=op['insurance'], column=col).value = f"=$D${inp['entry_ev']}*$D${inp['insurance']}"

        # EBITDA
        ws.cell(row=op['ebitda'], column=col).value = \
            f"={c}{op['revenue']}-{c}{op['om']}-{c}{op['land_lease']}-{c}{op['insurance']}"
        ws.cell(row=op['ebitda'], column=col).fill = GREEN_FILL
        ws.cell(row=op['ebitda'], column=col).font = BLACK_BOLD
        ws.cell(row=op['ebitda'], column=col).number_format = CURRENCY_FMT

        # Depreciation
        ws.cell(row=op['depreciation'], column=col).value = \
            f"=$D${inp['entry_ev']}*$D${inp['dep_basis']}/$D${inp['dep_life']}"

        # EBIT
        ws.cell(row=op['ebit'], column=col).value = f"={c}{op['ebitda']}-{c}{op['depreciation']}"

        # Pre-PTC Tax
        ws.cell(row=op['pre_ptc_tax'], column=col).value = f"=MAX(0,{c}{op['ebit']})*$D${inp['tax_rate']}"

        # Gross PTC (only during PTC period)
        ws.cell(row=op['gross_ptc'], column=col).value = \
            f"=IF({year}<=$D${inp['ptc_end_year']},{c}{op['generation']}*$D${inp['ptc_rate']}/1000000,0)"

        # PTC Applied = MIN(Gross PTC, Pre-PTC Tax) - CANNOT create negative taxes
        ws.cell(row=op['ptc_applied'], column=col).value = f"=MIN({c}{op['gross_ptc']},{c}{op['pre_ptc_tax']})"

        # Net Taxes
        ws.cell(row=op['taxes'], column=col).value = f"={c}{op['pre_ptc_tax']}-{c}{op['ptc_applied']}"

        # CFADS
        ws.cell(row=op['cfads'], column=col).value = f"={c}{op['ebitda']}-{c}{op['taxes']}"
        ws.cell(row=op['cfads'], column=col).fill = GREEN_FILL
        ws.cell(row=op['cfads'], column=col).font = BLACK_BOLD

    # ===== 3. SCULPTED DEBT =====
    print("[3] Fixing Sculpted Debt Schedule...")

    ws.cell(row=debt['beg_bal'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)

        # Beginning Balance
        if year == 1:
            ws.cell(row=debt['beg_bal'], column=col).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
        else:
            ws.cell(row=debt['beg_bal'], column=col).value = f"={prev}{debt['end_bal']}"

        # Interest
        ws.cell(row=debt['interest'], column=col).value = f"={c}{debt['beg_bal']}*$D${inp['int_rate']}"

        # Target DS = CFADS / Target DSCR
        ws.cell(row=debt['target_ds'], column=col).value = f"={c}{op['cfads']}/$D${inp['target_dscr']}"

        # Sculpted Principal = MAX(0, Target DS - Interest)
        ws.cell(row=debt['sculpted_prin'], column=col).value = f"=MAX(0,{c}{debt['target_ds']}-{c}{debt['interest']})"

        # Capped Principal = MIN(Sculpted, Beg Bal)
        ws.cell(row=debt['capped_prin'], column=col).value = f"=MIN({c}{debt['sculpted_prin']},{c}{debt['beg_bal']})"

        # Ending Balance
        ws.cell(row=debt['end_bal'], column=col).value = f"=MAX(0,{c}{debt['beg_bal']}-{c}{debt['capped_prin']})"

        # Actual DS
        ws.cell(row=debt['actual_ds'], column=col).value = f"={c}{debt['interest']}+{c}{debt['capped_prin']}"

        # DSCR
        ws.cell(row=debt['dscr'], column=col).value = f"=IF({c}{debt['actual_ds']}>0,{c}{op['cfads']}/{c}{debt['actual_ds']},0)"
        ws.cell(row=debt['dscr'], column=col).number_format = MULTIPLE_FMT

    # ===== 4. DSRA =====
    print("[4] Fixing DSRA...")

    ws.cell(row=dsra['beg'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)
        next_c = get_column_letter(col + 1) if col < 14 else c

        if col < 14:
            ws.cell(row=dsra['target'], column=col).value = f"={next_c}{debt['actual_ds']}*$D${inp['dsra_months']}/12"
        else:
            ws.cell(row=dsra['target'], column=col).value = 0

        if year == 1:
            ws.cell(row=dsra['beg'], column=col).value = f"=D{su['dsra_init']}"
        else:
            ws.cell(row=dsra['beg'], column=col).value = f"={prev}{dsra['end']}"

        ws.cell(row=dsra['funding'], column=col).value = f"={c}{dsra['target']}-{c}{dsra['beg']}"
        ws.cell(row=dsra['end'], column=col).value = f"={c}{dsra['beg']}+{c}{dsra['funding']}"

    # ===== 5. DISTRIBUTION WATERFALL =====
    print("[5] Fixing Distribution Waterfall...")

    for col in range(5, 15):
        c = get_column_letter(col)

        ws.cell(row=dist['avail'], column=col).value = \
            f"={c}{op['cfads']}-{c}{debt['actual_ds']}-{c}{dsra['funding']}"
        ws.cell(row=dist['distributable'], column=col).value = f"=MAX(0,{c}{dist['avail']})"

    # ===== 6. SOURCES & USES =====
    print("[6] Fixing Sources & Uses...")

    ws.cell(row=su['purchase'], column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=su['txn_fees'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['txn_fees']}"
    ws.cell(row=su['dsra_init'], column=4).value = f"=E{dsra['target']}"
    ws.cell(row=su['total_uses'], column=4).value = f"=D{su['purchase']}+D{su['txn_fees']}+D{su['dsra_init']}"

    ws.cell(row=su['debt'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
    ws.cell(row=su['equity'], column=4).value = f"=D{su['total_uses']}-D{su['debt']}"
    ws.cell(row=su['total_sources'], column=4).value = f"=D{su['debt']}+D{su['equity']}"
    ws.cell(row=su['balance'], column=4).value = f"=IF(ABS(D{su['total_uses']}-D{su['total_sources']})<0.01,\"PASS\",\"FAIL\")"

    # ===== 7. EXIT & RETURNS =====
    print("[7] Fixing Exit & Returns...")

    exit_year = 7
    exit_col = get_column_letter(4 + exit_year)

    ws.cell(row=exit_r['exit_ebitda'], column=4).value = f"={exit_col}{op['ebitda']}"
    ws.cell(row=exit_r['exit_ev'], column=4).value = f"=D{exit_r['exit_ebitda']}*$D${inp['exit_mult']}"
    ws.cell(row=exit_r['exit_debt'], column=4).value = f"={exit_col}{debt['end_bal']}"
    ws.cell(row=exit_r['exit_dsra'], column=4).value = f"={exit_col}{dsra['end']}"
    ws.cell(row=exit_r['exit_equity'], column=4).value = f"=D{exit_r['exit_ev']}-D{exit_r['exit_debt']}+D{exit_r['exit_dsra']}"

    ws.cell(row=exit_r['equity_cf'], column=4).value = f"=-D{su['equity']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}"
        elif year == exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}+D{exit_r['exit_equity']}"
        else:
            ws.cell(row=exit_r['equity_cf'], column=col).value = 0

    ws.cell(row=exit_r['irr'], column=4).value = f"=IRR(D{exit_r['equity_cf']}:N{exit_r['equity_cf']})"
    ws.cell(row=exit_r['irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['irr'], column=4).number_format = PERCENT_FMT

    ws.cell(row=exit_r['moic'], column=4).value = \
        f"=(SUM(E{exit_r['equity_cf']}:K{exit_r['equity_cf']}))/-D{exit_r['equity_cf']}"
    ws.cell(row=exit_r['moic'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['moic'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['moic'], column=4).number_format = MULTIPLE_FMT

    # ===== 8. AUDIT STRIP =====
    print("[8] Fixing Audit Strip...")

    ws.cell(row=3, column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=4, column=4).value = f"=E{op['ebitda']}"

    wb.save(filepath)
    print(f"\nSAVED: {filepath}")

    print("\n--- KEY FORMULAS ---")
    print(f"  E68 (Gen Y1): {ws.cell(row=68, column=5).value}")
    print(f"  E69 (Revenue): {ws.cell(row=69, column=5).value}")
    print(f"  E74 (EBITDA): {ws.cell(row=74, column=5).value}")
    print(f"  E81 (PTC Applied): {ws.cell(row=81, column=5).value}")
    print(f"  D111 (Purchase): {ws.cell(row=111, column=4).value}")

if __name__ == "__main__":
    fix_wind()
