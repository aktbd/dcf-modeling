#!/usr/bin/env python3
"""
Complete fix for Model_1_CCGT.xlsx - All sections with correct row references
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

# Style definitions
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'
CURRENCY_FMT = '#,##0.0_);(#,##0.0)'

def fix_ccgt():
    filepath = f"{DELIV_DIR}/Model_1_CCGT.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    # ===== VERIFIED INPUT ROW MAPPING =====
    inp = {
        'entry_ev': 28,       # 520
        'txn_fees': 29,       # 0.015
        'exit_mult': 30,      # 7
        'exit_year': 31,      # 7
        'capacity': 34,       # 365
        'heat_rate': 35,      # 7000
        'cap_factor': 36,     # 0.55
        'for_rate': 37,       # 0.05
        'power_price': 40,    # 48
        'cap_price': 41,      # 110
        'gas_price': 42,      # 3
        'escalation': 43,     # 0.025
        'vom': 46,            # 3.5
        'fom': 47,            # 15
        'leverage': 50,       # 0.45
        'int_rate': 51,       # 0.065
        'debt_tenor': 52,     # 7
        'dsra_months': 53,    # 6
        'sweep_pct': 54,      # 0.75
        'lockup': 55,         # 1.1
        'tax_rate': 58,       # 0.25
        'dep_basis': 59,      # 0.85
        'dep_life': 60,       # 15
    }

    # ===== CALCULATED ITEMS ROW MAPPING =====
    calc = {
        'ucap': 64,
        'generation': 65,
        'debt_amount': 66,
        'depreciation': 67,
        'sched_prin': 68,
    }

    # ===== OPERATING MODEL ROW MAPPING =====
    op = {
        'power_price': 72,
        'cap_price': 73,
        'gas_price': 74,
        'energy_rev': 76,
        'cap_rev': 77,
        'total_rev': 78,
        'fuel_cost': 80,
        'spark_spread': 81,
        'vom': 82,
        'fom': 83,
        'total_costs': 84,
        'ebitda': 86,
        'ebit': 90,
        'taxes': 91,
        'cfads': 92,
    }

    # ===== DEBT SCHEDULE ROW MAPPING =====
    debt = {
        'beg_bal': 96,
        'interest': 97,
        'sched_prin': 98,
        'ds_pre': 99,
        'excess': 100,
        'sweep': 101,
        'total_prin': 102,
        'end_bal': 103,
        'dscr': 104,
        'sched_ds': 106,  # For DSRA calculation (avoids circularity)
    }

    # ===== DSRA ROW MAPPING =====
    dsra = {
        'target': 108,
        'beg': 109,
        'funding': 110,
        'end': 111,
    }

    # ===== DISTRIBUTION WATERFALL =====
    dist = {
        'avail': 115,
        'lockup': 116,
        'distributable': 117,
    }

    # ===== SOURCES & USES =====
    su = {
        'purchase': 122,
        'txn_fees': 123,
        'dsra_init': 124,
        'total_uses': 125,
        'debt': 128,
        'equity': 129,
        'total_sources': 130,
        'balance': 132,
    }

    # ===== EXIT & RETURNS =====
    exit_r = {
        'exit_ebitda': 139,
        'exit_ev': 140,
        'exit_debt': 141,
        'exit_dsra': 142,
        'exit_equity': 143,
        'equity_cf': 146,
        'irr': 147,
        'moic': 148,
    }

    print("="*60)
    print("FIXING MODEL_1_CCGT.xlsx - COMPLETE REWRITE")
    print("="*60)

    # ===== 1. FIX CALCULATED ITEMS (Column D) =====
    print("\n[1] Fixing Calculated Items...")

    # UCAP = Capacity * (1 - FOR)
    ws.cell(row=calc['ucap'], column=4).value = f"=$D${inp['capacity']}*(1-$D${inp['for_rate']})"

    # Generation = Capacity * CF * 8760 * (1 - FOR)
    ws.cell(row=calc['generation'], column=4).value = f"=$D${inp['capacity']}*$D${inp['cap_factor']}*8760*(1-$D${inp['for_rate']})"

    # Debt Amount = Entry EV * Leverage
    ws.cell(row=calc['debt_amount'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"

    # Annual Depreciation = Entry EV * Dep Basis / Dep Life
    ws.cell(row=calc['depreciation'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['dep_basis']}/$D${inp['dep_life']}"

    # Scheduled Principal = Debt Amount / Debt Tenor
    ws.cell(row=calc['sched_prin'], column=4).value = f"=$D${calc['debt_amount']}/$D${inp['debt_tenor']}"

    # ===== 2. FIX OPERATING MODEL (Columns E-N) =====
    print("[2] Fixing Operating Model...")

    for col in range(5, 15):  # E=5 to N=14
        year = col - 4
        c = get_column_letter(col)

        # Price escalation
        ws.cell(row=op['power_price'], column=col).value = f"=$D${inp['power_price']}*(1+$D${inp['escalation']})^({year}-1)"
        ws.cell(row=op['cap_price'], column=col).value = f"=$D${inp['cap_price']}*(1+$D${inp['escalation']})^({year}-1)"
        ws.cell(row=op['gas_price'], column=col).value = f"=$D${inp['gas_price']}*(1+$D${inp['escalation']})^({year}-1)"

        # Revenue
        ws.cell(row=op['energy_rev'], column=col).value = f"=$D${calc['generation']}*{c}{op['power_price']}/1000000"
        ws.cell(row=op['cap_rev'], column=col).value = f"=$D${inp['capacity']}*{c}{op['cap_price']}*1000/1000000"
        ws.cell(row=op['total_rev'], column=col).value = f"={c}{op['energy_rev']}+{c}{op['cap_rev']}"

        # Costs
        ws.cell(row=op['fuel_cost'], column=col).value = f"=$D${calc['generation']}*$D${inp['heat_rate']}*{c}{op['gas_price']}/1000000000"
        ws.cell(row=op['spark_spread'], column=col).value = f"={c}{op['power_price']}-($D${inp['heat_rate']}/1000)*{c}{op['gas_price']}"
        ws.cell(row=op['vom'], column=col).value = f"=$D${calc['generation']}*$D${inp['vom']}/1000000"
        ws.cell(row=op['fom'], column=col).value = f"=$D${inp['capacity']}*$D${inp['fom']}*1000/1000000"
        ws.cell(row=op['total_costs'], column=col).value = f"={c}{op['fuel_cost']}+{c}{op['vom']}+{c}{op['fom']}"

        # EBITDA
        ws.cell(row=op['ebitda'], column=col).value = f"={c}{op['total_rev']}-{c}{op['total_costs']}"
        ws.cell(row=op['ebitda'], column=col).fill = GREEN_FILL
        ws.cell(row=op['ebitda'], column=col).font = BLACK_BOLD
        ws.cell(row=op['ebitda'], column=col).number_format = CURRENCY_FMT

        # EBIT = EBITDA - Depreciation
        ws.cell(row=op['ebit'], column=col).value = f"={c}{op['ebitda']}-$D${calc['depreciation']}"

        # Taxes = MAX(0, EBIT) * Tax Rate
        ws.cell(row=op['taxes'], column=col).value = f"=MAX(0,{c}{op['ebit']})*$D${inp['tax_rate']}"

        # CFADS = EBITDA - Taxes
        ws.cell(row=op['cfads'], column=col).value = f"={c}{op['ebitda']}-{c}{op['taxes']}"
        ws.cell(row=op['cfads'], column=col).fill = GREEN_FILL
        ws.cell(row=op['cfads'], column=col).font = BLACK_BOLD
        ws.cell(row=op['cfads'], column=col).number_format = CURRENCY_FMT

    # ===== 3. FIX DEBT SCHEDULE (Columns E-N) =====
    print("[3] Fixing Debt Schedule...")

    # Y0 column (D) - initialize
    ws.cell(row=debt['beg_bal'], column=4).value = 0
    ws.cell(row=debt['end_bal'], column=4).value = f"=$D${calc['debt_amount']}"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)

        # Beginning Balance: Y1 = Debt Amount, Y2+ = Prior End Balance
        if year == 1:
            ws.cell(row=debt['beg_bal'], column=col).value = f"=$D${calc['debt_amount']}"
        else:
            ws.cell(row=debt['beg_bal'], column=col).value = f"={prev}{debt['end_bal']}"

        # Interest = Beg Bal * Interest Rate
        ws.cell(row=debt['interest'], column=col).value = f"={c}{debt['beg_bal']}*$D${inp['int_rate']}"

        # Scheduled Principal = MIN(Sched Prin, Beg Bal)
        ws.cell(row=debt['sched_prin'], column=col).value = f"=MIN($D${calc['sched_prin']},{c}{debt['beg_bal']})"

        # Debt Service (Pre-Sweep) = Interest + Sched Principal
        ws.cell(row=debt['ds_pre'], column=col).value = f"={c}{debt['interest']}+{c}{debt['sched_prin']}"

        # Scheduled DS for DSRA (deterministic, no circularity)
        # = MAX(0, Debt - (Year-1)*Sched Prin) * Int Rate + Sched Prin
        ws.cell(row=debt['sched_ds'], column=col).value = \
            f"=MAX(0,$D${calc['debt_amount']}-({year}-1)*$D${calc['sched_prin']})*$D${inp['int_rate']}+$D${calc['sched_prin']}"

    # Now do Excess, Sweep, Total Prin, End Bal, DSCR (depend on DSRA funding)
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Excess Cash = MAX(0, CFADS - DS Pre - DSRA Funding)
        ws.cell(row=debt['excess'], column=col).value = f"=MAX(0,{c}{op['cfads']}-{c}{debt['ds_pre']}-{c}{dsra['funding']})"

        # Cash Sweep = MIN(Excess * Sweep%, Beg Bal - Sched Prin)
        ws.cell(row=debt['sweep'], column=col).value = \
            f"=MIN({c}{debt['excess']}*$D${inp['sweep_pct']},MAX(0,{c}{debt['beg_bal']}-{c}{debt['sched_prin']}))"

        # Total Principal = Sched + Sweep
        ws.cell(row=debt['total_prin'], column=col).value = f"={c}{debt['sched_prin']}+{c}{debt['sweep']}"

        # Ending Balance = MAX(0, Beg Bal - Total Principal)
        ws.cell(row=debt['end_bal'], column=col).value = f"=MAX(0,{c}{debt['beg_bal']}-{c}{debt['total_prin']})"

        # DSCR = IF(DS > 0, CFADS / DS, 0)
        ws.cell(row=debt['dscr'], column=col).value = f"=IF({c}{debt['ds_pre']}>0,{c}{op['cfads']}/{c}{debt['ds_pre']},0)"
        ws.cell(row=debt['dscr'], column=col).number_format = MULTIPLE_FMT

    # ===== 4. FIX DSRA (Columns D-N) =====
    print("[4] Fixing DSRA...")

    # Y0 (Column D) - DSRA initialized
    ws.cell(row=dsra['beg'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)
        next_c = get_column_letter(col + 1) if col < 14 else c

        # DSRA Target = Next Year's Sched DS * (DSRA Months / 12)
        # Final year target = 0 (released at exit)
        if col < 14:
            ws.cell(row=dsra['target'], column=col).value = f"={next_c}{debt['sched_ds']}*$D${inp['dsra_months']}/12"
        else:
            ws.cell(row=dsra['target'], column=col).value = 0

        # DSRA Beginning = Prior End (Y1 = D124 initial funding)
        if year == 1:
            ws.cell(row=dsra['beg'], column=col).value = f"=D{su['dsra_init']}"
        else:
            ws.cell(row=dsra['beg'], column=col).value = f"={prev}{dsra['end']}"

        # DSRA Funding = Target - Beginning (positive = funding, negative = release)
        ws.cell(row=dsra['funding'], column=col).value = f"={c}{dsra['target']}-{c}{dsra['beg']}"

        # DSRA Ending = Beginning + Funding
        ws.cell(row=dsra['end'], column=col).value = f"={c}{dsra['beg']}+{c}{dsra['funding']}"

    # ===== 5. FIX DISTRIBUTION WATERFALL (Columns E-N) =====
    print("[5] Fixing Distribution Waterfall...")

    for col in range(5, 15):
        c = get_column_letter(col)

        # Cash Available = CFADS - DS - Total Prin - DSRA Funding (but DS includes interest+sched prin)
        # Actually: Cash Avail = CFADS - Interest - Total Prin - DSRA Funding
        ws.cell(row=dist['avail'], column=col).value = \
            f"={c}{op['cfads']}-{c}{debt['interest']}-{c}{debt['total_prin']}-{c}{dsra['funding']}"

        # Lock-up Test = IF(DSCR >= Lockup, "PASS", "LOCKED")
        ws.cell(row=dist['lockup'], column=col).value = f"=IF({c}{debt['dscr']}>=$D${inp['lockup']},\"PASS\",\"LOCKED\")"

        # Distributable = IF(PASS, MAX(0, Cash Avail), 0)
        ws.cell(row=dist['distributable'], column=col).value = \
            f"=IF({c}{dist['lockup']}=\"PASS\",MAX(0,{c}{dist['avail']}),0)"

    # ===== 6. FIX SOURCES & USES (Column D only) =====
    print("[6] Fixing Sources & Uses...")

    # Purchase Price = Entry EV
    ws.cell(row=su['purchase'], column=4).value = f"=$D${inp['entry_ev']}"

    # Transaction Fees = Entry EV * Txn Fee %
    ws.cell(row=su['txn_fees'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['txn_fees']}"

    # DSRA Initial = Y1 DSRA Target
    ws.cell(row=su['dsra_init'], column=4).value = f"=E{dsra['target']}"

    # Total Uses = Purchase + Fees + DSRA
    ws.cell(row=su['total_uses'], column=4).value = f"=D{su['purchase']}+D{su['txn_fees']}+D{su['dsra_init']}"

    # Funded Debt = Entry EV * Leverage
    ws.cell(row=su['debt'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"

    # Entry Equity = Total Uses - Debt (plug)
    ws.cell(row=su['equity'], column=4).value = f"=D{su['total_uses']}-D{su['debt']}"

    # Total Sources = Debt + Equity
    ws.cell(row=su['total_sources'], column=4).value = f"=D{su['debt']}+D{su['equity']}"

    # Balance Check = Uses - Sources (should be 0)
    ws.cell(row=su['balance'], column=4).value = f"=D{su['total_uses']}-D{su['total_sources']}"

    # ===== 7. FIX EXIT & RETURNS =====
    print("[7] Fixing Exit & Returns...")

    exit_year = 7
    exit_col = get_column_letter(4 + exit_year)  # K for Y7

    # Exit EBITDA = EBITDA in exit year
    ws.cell(row=exit_r['exit_ebitda'], column=4).value = f"={exit_col}{op['ebitda']}"

    # Exit EV = Exit EBITDA * Exit Multiple
    ws.cell(row=exit_r['exit_ev'], column=4).value = f"=D{exit_r['exit_ebitda']}*$D${inp['exit_mult']}"

    # Exit Debt = End Balance at exit
    ws.cell(row=exit_r['exit_debt'], column=4).value = f"={exit_col}{debt['end_bal']}"

    # DSRA Release = DSRA End at exit
    ws.cell(row=exit_r['exit_dsra'], column=4).value = f"={exit_col}{dsra['end']}"

    # Exit Equity = Exit EV - Debt + DSRA
    ws.cell(row=exit_r['exit_equity'], column=4).value = f"=D{exit_r['exit_ev']}-D{exit_r['exit_debt']}+D{exit_r['exit_dsra']}"

    # ===== 8. FIX EQUITY CASH FLOW FOR IRR =====
    print("[8] Fixing Equity Cash Flow & Returns...")

    # Y0 = -Entry Equity (negative outflow)
    ws.cell(row=exit_r['equity_cf'], column=4).value = f"=-D{su['equity']}"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        if year < exit_year:
            # Pre-exit: Distributable cash
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}"
        elif year == exit_year:
            # Exit year: Distributable + Exit Equity Proceeds
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}+D{exit_r['exit_equity']}"
        else:
            # Post-exit: 0
            ws.cell(row=exit_r['equity_cf'], column=col).value = 0

    # Levered IRR
    ws.cell(row=exit_r['irr'], column=4).value = f"=IRR(D{exit_r['equity_cf']}:N{exit_r['equity_cf']})"
    ws.cell(row=exit_r['irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['irr'], column=4).number_format = PERCENT_FMT

    # MOIC = Total Distributions / Entry Equity
    ws.cell(row=exit_r['moic'], column=4).value = \
        f"=(SUM(E{exit_r['equity_cf']}:K{exit_r['equity_cf']}))/-D{exit_r['equity_cf']}"
    ws.cell(row=exit_r['moic'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['moic'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['moic'], column=4).number_format = MULTIPLE_FMT

    # ===== 9. FIX AUDIT STRIP =====
    print("[9] Fixing Audit Strip...")

    ws.cell(row=3, column=4).value = f"=$D${inp['entry_ev']}"  # Entry EV
    ws.cell(row=4, column=4).value = f"=E{op['ebitda']}"  # Y1 EBITDA

    # Fix QA checks
    ws.cell(row=22, column=4).value = f"=D{su['balance']}"  # S&U Balance
    ws.cell(row=23, column=4).value = f"=IF({exit_col}{debt['end_bal']}<1,\"PASS\",\"FAIL\")"  # Debt Payoff
    ws.cell(row=24, column=4).value = f"=IF(MIN(E{debt['dscr']}:K{debt['dscr']})>=1.25,\"PASS\",\"FAIL\")"  # Min DSCR

    # ===== SAVE =====
    wb.save(filepath)
    print(f"\n{'='*60}")
    print(f"SAVED: {filepath}")
    print('='*60)

    return filepath

def validate_ccgt(filepath):
    """Validate the fixed model by checking key outputs"""
    wb = load_workbook(filepath, data_only=True)
    ws = wb['Model_Standard']

    print("\n" + "="*60)
    print("VALIDATION - Calculated Values")
    print("="*60)

    # Read calculated values
    checks = {
        'Entry EV (D28)': ws.cell(row=28, column=4).value,
        'UCAP (D64)': ws.cell(row=64, column=4).value,
        'Generation (D65)': ws.cell(row=65, column=4).value,
        'Debt Amount (D66)': ws.cell(row=66, column=4).value,
        'Y1 Power Price (E72)': ws.cell(row=72, column=5).value,
        'Y1 Energy Rev (E76)': ws.cell(row=76, column=5).value,
        'Y1 Fuel Cost (E80)': ws.cell(row=80, column=5).value,
        'Y1 EBITDA (E86)': ws.cell(row=86, column=5).value,
        'Y1 CFADS (E92)': ws.cell(row=92, column=5).value,
        'Y1 Beg Bal (E96)': ws.cell(row=96, column=5).value,
        'Y1 Interest (E97)': ws.cell(row=97, column=5).value,
        'Y1 DSCR (E104)': ws.cell(row=104, column=5).value,
        'S&U Balance (D132)': ws.cell(row=132, column=4).value,
        'IRR (D147)': ws.cell(row=147, column=4).value,
        'MOIC (D148)': ws.cell(row=148, column=4).value,
    }

    for name, val in checks.items():
        print(f"  {name}: {val}")

    # Sanity checks
    print("\n" + "-"*40)
    print("SANITY CHECKS:")

    issues = []

    ebitda = checks['Y1 EBITDA (E86)']
    if ebitda and (ebitda < 0 or ebitda > 200):
        issues.append(f"  EBITDA {ebitda} outside expected range (0-200)")

    dscr = checks['Y1 DSCR (E104)']
    if dscr and (dscr < 1.0 or dscr > 5.0):
        issues.append(f"  DSCR {dscr} outside expected range (1.0-5.0)")

    irr = checks['IRR (D147)']
    if irr and (irr < 0.05 or irr > 0.40):
        issues.append(f"  IRR {irr:.1%} outside expected range (5%-40%)")

    balance = checks['S&U Balance (D132)']
    if balance and abs(balance) > 0.01:
        issues.append(f"  S&U Balance {balance} should be ~0")

    if issues:
        print("ISSUES FOUND:")
        for issue in issues:
            print(issue)
    else:
        print("  All checks PASSED!")

    return len(issues) == 0


if __name__ == "__main__":
    filepath = fix_ccgt()

    # Need to reopen with data_only=True to see calculated values
    # But formulas need to be calculated first (requires Excel or xlcalc)
    print("\nNOTE: Open file in Excel to verify calculated values.")
    print("Formula checks:")

    wb = load_workbook(filepath, data_only=False)
    ws = wb['Model_Standard']

    print(f"  E72 (Power Price): {ws.cell(row=72, column=5).value}")
    print(f"  E76 (Energy Rev): {ws.cell(row=76, column=5).value}")
    print(f"  E80 (Fuel Cost): {ws.cell(row=80, column=5).value}")
    print(f"  E86 (EBITDA): {ws.cell(row=86, column=5).value}")
    print(f"  E97 (Interest): {ws.cell(row=97, column=5).value}")
    print(f"  E104 (DSCR): {ws.cell(row=104, column=5).value}")
    print(f"  D122 (Purchase): {ws.cell(row=122, column=4).value}")
    print(f"  D132 (Balance): {ws.cell(row=132, column=4).value}")
    print(f"  D147 (IRR): {ws.cell(row=147, column=4).value}")
