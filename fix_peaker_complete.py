#!/usr/bin/env python3
"""
Complete fix for Model_2_Peaker.xlsx
Key difference from CCGT: Generation = Capacity × Dispatch Hours × Availability
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'
CURRENCY_FMT = '#,##0.0_);(#,##0.0)'

def fix_peaker():
    filepath = f"{DELIV_DIR}/Model_2_Peaker.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    # ===== INPUT ROW MAPPING =====
    inp = {
        'entry_ev': 28,       # 110
        'txn_fees': 29,       # 0.015
        'exit_mult': 30,      # 7
        'exit_year': 31,      # 7
        'capacity': 34,       # 200
        'heat_rate': 35,      # 10500
        'availability': 36,   # 0.94
        'for_rate': 37,       # 0.05
        'dispatch_hrs': 38,   # 400
        'power_price': 41,    # 95
        'cap_price': 42,      # 60
        'gas_price': 43,      # 3.25
        'escalation': 44,     # 0.025
        'vom': 47,            # 4
        'fom': 48,            # 12
        'mm_year': 49,        # 5
        'mm_cost': 50,        # 4
        'leverage': 53,       # 0.3
        'int_rate': 54,       # 0.07
        'debt_tenor': 55,     # 7
        'dsra_months': 56,    # 6
        'sweep_pct': 57,      # 0.5
        'lockup': 58,         # 1.1
        'tax_rate': 61,       # 0.25
        'dep_basis': 62,      # 0.85
        'dep_life': 63,       # 15
    }

    calc = {
        'ucap': 66,
        'generation': 67,
        'debt_amount': 68,
        'depreciation': 69,
        'sched_prin': 70,
    }

    op = {
        'power_price': 73,
        'cap_price': 74,
        'gas_price': 75,
        'energy_rev': 77,
        'cap_rev': 78,
        'total_rev': 79,
        'fuel_cost': 81,
        'spark_spread': 82,
        'vom': 83,
        'fom': 84,
        'total_costs': 85,
        'ebitda': 87,
        'ebit': 90,
        'taxes': 91,
        'mm': 92,  # Major Maintenance
        'cfads': 93,
    }

    debt = {
        'beg_bal': 96,
        'interest': 97,
        'sched_prin': 98,
        'ds_pre': 99,
        'excess': 100,
        'sweep': 101,
        'total_prin': 102,
        'end_bal': 103,
        'dscr': 104,
        'sched_ds': 105,
    }

    dsra = {
        'target': 107,
        'beg': 108,
        'funding': 109,
        'end': 110,
    }

    dist = {
        'avail': 113,
        'lockup': 114,
        'distributable': 115,
    }

    su = {
        'purchase': 119,
        'txn_fees': 120,
        'dsra_init': 121,
        'total_uses': 122,
        'debt': 125,
        'equity': 126,
        'total_sources': 127,
        'balance': 128,
    }

    exit_r = {
        'exit_ev': 131,
        'exit_debt': 132,
        'exit_dsra': 133,
        'exit_equity': 134,
        'equity_cf': 137,
        'irr': 138,
        'moic': 139,
        'unlev_cf': 142,
        'unlev_irr': 143,
    }

    print("="*60)
    print("FIXING MODEL_2_PEAKER.xlsx")
    print("="*60)

    # ===== 1. CALCULATED ITEMS =====
    print("[1] Fixing Calculated Items...")

    # UCAP = Capacity * (1 - FOR)
    ws.cell(row=calc['ucap'], column=4).value = f"=$D${inp['capacity']}*(1-$D${inp['for_rate']})"

    # PEAKER: Generation = Capacity × Dispatch Hours × Availability
    ws.cell(row=calc['generation'], column=4).value = f"=$D${inp['capacity']}*$D${inp['dispatch_hrs']}*$D${inp['availability']}"

    # Debt Amount = Entry EV * Leverage
    ws.cell(row=calc['debt_amount'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"

    # Annual Depreciation
    ws.cell(row=calc['depreciation'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['dep_basis']}/$D${inp['dep_life']}"

    # Scheduled Principal
    ws.cell(row=calc['sched_prin'], column=4).value = f"=$D${calc['debt_amount']}/$D${inp['debt_tenor']}"

    # ===== 2. OPERATING MODEL =====
    print("[2] Fixing Operating Model...")

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Price escalation
        ws.cell(row=op['power_price'], column=col).value = f"=$D${inp['power_price']}*(1+$D${inp['escalation']})^({year}-1)"
        ws.cell(row=op['cap_price'], column=col).value = f"=$D${inp['cap_price']}*(1+$D${inp['escalation']})^({year}-1)"
        ws.cell(row=op['gas_price'], column=col).value = f"=$D${inp['gas_price']}*(1+$D${inp['escalation']})^({year}-1)"

        # Revenue (Peakers: capacity revenue dominates)
        ws.cell(row=op['energy_rev'], column=col).value = f"=$D${calc['generation']}*{c}{op['power_price']}/1000000"
        ws.cell(row=op['cap_rev'], column=col).value = f"=$D${inp['capacity']}*{c}{op['cap_price']}*1000/1000000"
        ws.cell(row=op['total_rev'], column=col).value = f"={c}{op['energy_rev']}+{c}{op['cap_rev']}"

        # Costs
        ws.cell(row=op['fuel_cost'], column=col).value = f"=$D${calc['generation']}*$D${inp['heat_rate']}*{c}{op['gas_price']}/1000000000"
        ws.cell(row=op['spark_spread'], column=col).value = f"={c}{op['power_price']}-($D${inp['heat_rate']}/1000)*{c}{op['gas_price']}"
        ws.cell(row=op['vom'], column=col).value = f"=$D${calc['generation']}*$D${inp['vom']}/1000000"
        ws.cell(row=op['fom'], column=col).value = f"=$D${inp['capacity']}*$D${inp['fom']}*1000/1000000"
        ws.cell(row=op['total_costs'], column=col).value = f"={c}{op['fuel_cost']}+{c}{op['vom']}+{c}{op['fom']}"

        # EBITDA
        ws.cell(row=op['ebitda'], column=col).value = f"={c}{op['total_rev']}-{c}{op['total_costs']}"
        ws.cell(row=op['ebitda'], column=col).fill = GREEN_FILL
        ws.cell(row=op['ebitda'], column=col).font = BLACK_BOLD
        ws.cell(row=op['ebitda'], column=col).number_format = CURRENCY_FMT

        # EBIT
        ws.cell(row=op['ebit'], column=col).value = f"={c}{op['ebitda']}-$D${calc['depreciation']}"

        # Taxes
        ws.cell(row=op['taxes'], column=col).value = f"=MAX(0,{c}{op['ebit']})*$D${inp['tax_rate']}"

        # Major Maintenance (Year 5 only)
        ws.cell(row=op['mm'], column=col).value = f"=IF({year}=$D${inp['mm_year']},$D${inp['mm_cost']},0)"

        # CFADS = EBITDA - Taxes - Major Maintenance
        ws.cell(row=op['cfads'], column=col).value = f"={c}{op['ebitda']}-{c}{op['taxes']}-{c}{op['mm']}"
        ws.cell(row=op['cfads'], column=col).fill = GREEN_FILL
        ws.cell(row=op['cfads'], column=col).font = BLACK_BOLD

    # ===== 3. DEBT SCHEDULE =====
    print("[3] Fixing Debt Schedule...")

    ws.cell(row=debt['beg_bal'], column=4).value = 0
    ws.cell(row=debt['end_bal'], column=4).value = f"=$D${calc['debt_amount']}"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)

        if year == 1:
            ws.cell(row=debt['beg_bal'], column=col).value = f"=$D${calc['debt_amount']}"
        else:
            ws.cell(row=debt['beg_bal'], column=col).value = f"={prev}{debt['end_bal']}"

        ws.cell(row=debt['interest'], column=col).value = f"={c}{debt['beg_bal']}*$D${inp['int_rate']}"
        ws.cell(row=debt['sched_prin'], column=col).value = f"=MIN($D${calc['sched_prin']},{c}{debt['beg_bal']})"
        ws.cell(row=debt['ds_pre'], column=col).value = f"={c}{debt['interest']}+{c}{debt['sched_prin']}"
        ws.cell(row=debt['sched_ds'], column=col).value = \
            f"=MAX(0,$D${calc['debt_amount']}-({year}-1)*$D${calc['sched_prin']})*$D${inp['int_rate']}+$D${calc['sched_prin']}"

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        ws.cell(row=debt['excess'], column=col).value = f"=MAX(0,{c}{op['cfads']}-{c}{debt['ds_pre']}-{c}{dsra['funding']})"
        ws.cell(row=debt['sweep'], column=col).value = \
            f"=MIN({c}{debt['excess']}*$D${inp['sweep_pct']},MAX(0,{c}{debt['beg_bal']}-{c}{debt['sched_prin']}))"
        ws.cell(row=debt['total_prin'], column=col).value = f"={c}{debt['sched_prin']}+{c}{debt['sweep']}"
        ws.cell(row=debt['end_bal'], column=col).value = f"=MAX(0,{c}{debt['beg_bal']}-{c}{debt['total_prin']})"
        ws.cell(row=debt['dscr'], column=col).value = f"=IF({c}{debt['ds_pre']}>0,{c}{op['cfads']}/{c}{debt['ds_pre']},0)"
        ws.cell(row=debt['dscr'], column=col).number_format = MULTIPLE_FMT

    # ===== 4. DSRA =====
    print("[4] Fixing DSRA...")

    ws.cell(row=dsra['beg'], column=4).value = 0

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        prev = get_column_letter(col - 1)
        next_c = get_column_letter(col + 1) if col < 14 else c

        if col < 14:
            ws.cell(row=dsra['target'], column=col).value = f"={next_c}{debt['sched_ds']}*$D${inp['dsra_months']}/12"
        else:
            ws.cell(row=dsra['target'], column=col).value = 0

        if year == 1:
            ws.cell(row=dsra['beg'], column=col).value = f"=D{su['dsra_init']}"
        else:
            ws.cell(row=dsra['beg'], column=col).value = f"={prev}{dsra['end']}"

        ws.cell(row=dsra['funding'], column=col).value = f"={c}{dsra['target']}-{c}{dsra['beg']}"
        ws.cell(row=dsra['end'], column=col).value = f"={c}{dsra['beg']}+{c}{dsra['funding']}"

    # ===== 5. DISTRIBUTION WATERFALL =====
    print("[5] Fixing Distribution Waterfall...")

    for col in range(5, 15):
        c = get_column_letter(col)

        ws.cell(row=dist['avail'], column=col).value = \
            f"={c}{op['cfads']}-{c}{debt['interest']}-{c}{debt['total_prin']}-{c}{dsra['funding']}"
        ws.cell(row=dist['lockup'], column=col).value = f"=IF({c}{debt['dscr']}>=$D${inp['lockup']},\"PASS\",\"LOCKED\")"
        ws.cell(row=dist['distributable'], column=col).value = \
            f"=IF({c}{dist['lockup']}=\"PASS\",MAX(0,{c}{dist['avail']}),0)"

    # ===== 6. SOURCES & USES =====
    print("[6] Fixing Sources & Uses...")

    ws.cell(row=su['purchase'], column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=su['txn_fees'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['txn_fees']}"
    ws.cell(row=su['dsra_init'], column=4).value = f"=E{dsra['target']}"
    ws.cell(row=su['total_uses'], column=4).value = f"=D{su['purchase']}+D{su['txn_fees']}+D{su['dsra_init']}"
    ws.cell(row=su['debt'], column=4).value = f"=$D${inp['entry_ev']}*$D${inp['leverage']}"
    ws.cell(row=su['equity'], column=4).value = f"=D{su['total_uses']}-D{su['debt']}"
    ws.cell(row=su['total_sources'], column=4).value = f"=D{su['debt']}+D{su['equity']}"
    ws.cell(row=su['balance'], column=4).value = f"=IF(ABS(D{su['total_uses']}-D{su['total_sources']})<0.01,\"PASS\",\"FAIL\")"

    # ===== 7. EXIT & RETURNS =====
    print("[7] Fixing Exit & Returns...")

    exit_year = 7
    exit_col = get_column_letter(4 + exit_year)

    ws.cell(row=exit_r['exit_ev'], column=4).value = f"={exit_col}{op['ebitda']}*$D${inp['exit_mult']}"
    ws.cell(row=exit_r['exit_debt'], column=4).value = f"={exit_col}{debt['end_bal']}"
    ws.cell(row=exit_r['exit_dsra'], column=4).value = f"={exit_col}{dsra['end']}"
    ws.cell(row=exit_r['exit_equity'], column=4).value = f"=D{exit_r['exit_ev']}-D{exit_r['exit_debt']}+D{exit_r['exit_dsra']}"

    # Equity CF
    ws.cell(row=exit_r['equity_cf'], column=4).value = f"=-D{su['equity']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}"
        elif year == exit_year:
            ws.cell(row=exit_r['equity_cf'], column=col).value = f"={c}{dist['distributable']}+D{exit_r['exit_equity']}"
        else:
            ws.cell(row=exit_r['equity_cf'], column=col).value = 0

    ws.cell(row=exit_r['irr'], column=4).value = f"=IRR(D{exit_r['equity_cf']}:N{exit_r['equity_cf']})"
    ws.cell(row=exit_r['irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['irr'], column=4).number_format = PERCENT_FMT

    ws.cell(row=exit_r['moic'], column=4).value = \
        f"=(SUM(E{exit_r['equity_cf']}:K{exit_r['equity_cf']}))/-D{exit_r['equity_cf']}"
    ws.cell(row=exit_r['moic'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['moic'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['moic'], column=4).number_format = MULTIPLE_FMT

    # Unlevered CF & IRR
    ws.cell(row=exit_r['unlev_cf'], column=4).value = f"=-$D${inp['entry_ev']}-D{su['txn_fees']}"
    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)
        if year < exit_year:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = f"={c}{op['cfads']}"
        elif year == exit_year:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = f"={c}{op['cfads']}+D{exit_r['exit_ev']}"
        else:
            ws.cell(row=exit_r['unlev_cf'], column=col).value = 0

    ws.cell(row=exit_r['unlev_irr'], column=4).value = f"=IRR(D{exit_r['unlev_cf']}:N{exit_r['unlev_cf']})"
    ws.cell(row=exit_r['unlev_irr'], column=4).fill = GREEN_FILL
    ws.cell(row=exit_r['unlev_irr'], column=4).font = BLACK_BOLD
    ws.cell(row=exit_r['unlev_irr'], column=4).number_format = PERCENT_FMT

    # ===== 8. AUDIT STRIP =====
    print("[8] Fixing Audit Strip...")

    ws.cell(row=3, column=4).value = f"=$D${inp['entry_ev']}"
    ws.cell(row=4, column=4).value = f"=E{op['ebitda']}"

    wb.save(filepath)
    print(f"\nSAVED: {filepath}")

    # Validation
    print("\n--- KEY FORMULAS ---")
    print(f"  D67 (Generation): {ws.cell(row=67, column=4).value}")
    print(f"  E77 (Energy Rev): {ws.cell(row=77, column=5).value}")
    print(f"  E87 (EBITDA): {ws.cell(row=87, column=5).value}")
    print(f"  D119 (Purchase): {ws.cell(row=119, column=4).value}")

if __name__ == "__main__":
    fix_peaker()
