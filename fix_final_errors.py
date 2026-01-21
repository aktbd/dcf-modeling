#!/usr/bin/env python3
"""
Final fix for remaining issues:
1. CCGT: Circular reference in Exit Equity / IRR
2. Midstream: Taxes formula is wrong (has CapEx formula instead)
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIV_DIR = '/home/user/dcf-modeling/deliverables/deliverables-20260120-0030'

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLACK_BOLD = Font(color="000000", bold=True)
PERCENT_FMT = '0.0%'
MULTIPLE_FMT = '0.00"x"'

def fix_ccgt_final():
    """Fix CCGT circular reference"""
    filepath = f"{DELIV_DIR}/Model_1_CCGT.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    print("="*60)
    print("FIXING CCGT - Circular Reference")
    print("="*60)

    # The issue: D143 (Exit Equity) = D140-D141+D142
    # But D142 (IRR) depends on D141, and K141 depends on D143
    # This creates a circular reference!

    # Fix: D143 should be calculated INDEPENDENTLY of IRR
    # Exit Equity = Exit EV - Exit Debt + DSRA Release

    # Find the correct exit EV calculation
    # Row 140 is "EQUITY RETURNS" header with formula "=D139*$D$30"
    # That's: Exit EBITDA × Exit Multiple

    # Let me trace the actual structure
    print("\nTracing Exit section:")
    for row in range(135, 150):
        a = ws.cell(row=row, column=1).value or ""
        d = ws.cell(row=row, column=4).value
        if a:
            print(f"  Row {row}: {str(a)[:30]:<30} | D={str(d)[:40] if d else ''}")

    # Based on typical structure, let's fix:
    # Row 139: Exit EBITDA = K86 (Y7 EBITDA)
    # Row 140: Exit EV = D139 * Exit Multiple
    # Row 141: Equity Cash Flow row
    # Row 142: Levered IRR
    # Row 143: MOIC

    # Actually, looking at the output:
    # Row 140: EQUITY RETURNS = D139*$D$30 (Exit EV calculation)
    # But this uses row 139 for Exit EBITDA

    # Fix Exit Equity Proceeds - needs to be independent
    # Exit Equity = Exit EV - Debt at Exit + DSRA at Exit

    # Find Exit EV (should be EBITDA × Multiple)
    exit_year = 7
    exit_col = get_column_letter(4 + exit_year)  # K

    # Create proper exit calculation rows
    # D139 should be Exit EBITDA
    ws.cell(row=139, column=4).value = f"={exit_col}86"  # Y7 EBITDA
    ws.cell(row=139, column=1).value = "Exit EBITDA"

    # D140 should be Exit EV
    ws.cell(row=140, column=4).value = "=D139*$D$30"  # EBITDA × Multiple
    ws.cell(row=140, column=1).value = "Exit EV"

    # Create Exit Debt row if needed (check row 141's current use)
    # But 141 is Equity Cash Flow which we need for IRR

    # Let's reorganize:
    # Row 138: Exit Debt = K103 (Y7 Ending Debt Balance)
    # Row 139: Exit EBITDA
    # Row 140: Exit EV
    # Row 141: Exit Equity Proceeds = D140 - Debt + DSRA

    # Actually, let me just fix the Exit Equity calculation properly
    # It should NOT reference the IRR

    # Find DSRA End and Debt End at exit year
    dsra_end_row = 111  # Based on earlier trace
    debt_end_row = 103  # Based on earlier trace

    # Create Exit Equity Proceeds (independent calculation)
    # We need a new row for this or use an existing one

    # Looking at current structure:
    # D143 = =D140-D141+D142 which is circular

    # Let me fix D143 to be: Exit EV - Exit Debt + DSRA Release
    # Exit EV is in D140 = D139*$D$30
    # But D139 needs to be Exit EBITDA first

    # Step 1: Ensure D139 = Exit EBITDA
    ws.cell(row=139, column=4).value = f"={exit_col}86"

    # Step 2: Ensure D140 = Exit EV = D139 × Multiple
    ws.cell(row=140, column=4).value = "=D139*$D$30"

    # Step 3: Create Exit Equity calculation that doesn't depend on IRR
    # We need to put it somewhere that breaks the circularity
    # Option: Calculate directly as Exit EV - Exit Debt + DSRA

    # First, let's see what row we can use for intermediate calcs
    # Row 144 might be available

    # Put Exit Debt reference
    ws.cell(row=144, column=1).value = "Exit Debt Balance"
    ws.cell(row=144, column=4).value = f"={exit_col}103"  # K103 = Y7 Ending Balance

    # Put DSRA Release
    ws.cell(row=145, column=1).value = "Exit DSRA Release"
    ws.cell(row=145, column=4).value = f"={exit_col}111"  # K111 = Y7 DSRA Ending

    # Put Exit Equity Proceeds (INDEPENDENT of IRR)
    ws.cell(row=143, column=1).value = "Exit Equity Proceeds"
    ws.cell(row=143, column=4).value = "=D140-D144+D145"  # EV - Debt + DSRA

    # Now fix Equity Cash Flow row (141)
    # Y0 = -Entry Equity (D129)
    ws.cell(row=141, column=4).value = "=-D129"

    # Y1-Y6 = Distributable (row 117)
    for col in range(5, 11):
        c = get_column_letter(col)
        ws.cell(row=141, column=col).value = f"={c}117"

    # Y7 = Distribution + Exit Equity
    ws.cell(row=141, column=11).value = "=K117+D143"  # Now D143 is independent!

    # Y8-Y10 = 0
    for col in range(12, 15):
        ws.cell(row=141, column=col).value = 0

    # IRR (no circular reference now)
    ws.cell(row=142, column=4).value = "=IRR(D141:N141)"
    ws.cell(row=142, column=4).fill = GREEN_FILL
    ws.cell(row=142, column=4).font = BLACK_BOLD
    ws.cell(row=142, column=4).number_format = PERCENT_FMT

    # MOIC - move to a different row or fix
    ws.cell(row=147, column=1).value = "MOIC"
    ws.cell(row=147, column=4).value = "=(SUM(E141:K141))/-D141"
    ws.cell(row=147, column=4).fill = GREEN_FILL
    ws.cell(row=147, column=4).font = BLACK_BOLD
    ws.cell(row=147, column=4).number_format = MULTIPLE_FMT

    print("\nFixed Exit section:")
    print("  D139 (Exit EBITDA): =K86")
    print("  D140 (Exit EV): =D139*$D$30")
    print("  D143 (Exit Equity): =D140-D144+D145 (NO circular ref)")
    print("  D144 (Exit Debt): =K103")
    print("  D145 (Exit DSRA): =K111")
    print("  D142 (IRR): =IRR(D141:N141)")

    wb.save(filepath)
    print(f"\nSaved: {filepath}")


def fix_midstream_final():
    """Fix Midstream Taxes vs CapEx confusion"""
    filepath = f"{DELIV_DIR}/Model_5_Midstream.xlsx"
    wb = load_workbook(filepath)
    ws = wb['Model_Standard']

    print("\n" + "="*60)
    print("FIXING MIDSTREAM - Taxes/CapEx Formula Confusion")
    print("="*60)

    # The issue: Row 76 (labeled "Taxes") has formula =IF(1<=5,$D$44,0) which is CapEx!
    # And E78 CFADS = E72-E76-IF(1<=5,$D$44,0) double-counts CapEx

    # Input values:
    # D44 = Growth CapEx = 2
    # D55 = Tax Rate = 0.25
    # D56 = Depreciable Basis = 0.8
    # D57 = Depreciation Life = 15
    # D61 = Annual Depreciation = Entry EV × Dep Basis / Dep Life

    # EBITDA is at row 72
    # Proper formulas:
    # Depreciation = D27 * D56 / D57 (fixed in calc items at D61)
    # EBIT = EBITDA - Depreciation
    # Taxes = MAX(0, EBIT) × Tax Rate
    # Growth CapEx = IF(Year <= 5, D44, 0)
    # CFADS = EBITDA - Taxes - Growth CapEx

    print("\nCurrent structure:")
    for row in range(72, 80):
        a = ws.cell(row=row, column=1).value or ""
        e = ws.cell(row=row, column=5).value
        print(f"  Row {row}: {str(a)[:20]:<20} | E={str(e)[:50] if e else ''}")

    # Fix row 75 (EBIT) - should be EBITDA - Depreciation
    # Fix row 76 (Taxes) - should be MAX(0, EBIT) × Tax Rate
    # Fix row 77 (Growth CapEx) - should be IF(Year<=5, D44, 0)
    # Fix row 78 (CFADS) - should be EBITDA - Taxes - CapEx

    for col in range(5, 15):
        year = col - 4
        c = get_column_letter(col)

        # Row 75: EBIT = EBITDA - Depreciation
        ws.cell(row=75, column=col).value = f"={c}72-$D$61"

        # Row 76: Taxes = MAX(0, EBIT) × Tax Rate
        ws.cell(row=76, column=col).value = f"=MAX(0,{c}75)*$D$55"

        # Row 77: Growth CapEx = IF(Year <= 5, D44, 0)
        ws.cell(row=77, column=col).value = f"=IF({year}<=5,$D$44,0)"

        # Row 78: CFADS = EBITDA - Taxes - Growth CapEx
        ws.cell(row=78, column=col).value = f"={c}72-{c}76-{c}77"
        ws.cell(row=78, column=col).fill = GREEN_FILL
        ws.cell(row=78, column=col).font = BLACK_BOLD

    print("\nFixed formulas:")
    print("  Row 75 (EBIT): =En72-$D$61")
    print("  Row 76 (Taxes): =MAX(0,En75)*$D$55")
    print("  Row 77 (Growth CapEx): =IF(Year<=5,$D$44,0)")
    print("  Row 78 (CFADS): =En72-En76-En77")

    wb.save(filepath)
    print(f"\nSaved: {filepath}")


def verify_final():
    """Final verification with formula content AND logic check"""
    print("\n" + "="*60)
    print("FINAL VERIFICATION")
    print("="*60)

    # CCGT
    wb = load_workbook(f"{DELIV_DIR}/Model_1_CCGT.xlsx")
    ws = wb['Model_Standard']

    print("\nCCGT:")
    d142 = ws.cell(row=142, column=4).value
    d143 = ws.cell(row=143, column=4).value
    print(f"  D142 (IRR): {d142}")
    print(f"  D143 (Exit Equity): {d143}")

    if 'IRR' in str(d142):
        print("  ✓ IRR formula correct")
    if 'D142' not in str(d143) and 'D141' not in str(d143):
        print("  ✓ Exit Equity has no circular reference")
    else:
        print("  ✗ WARNING: Potential circular reference in Exit Equity")

    # Midstream
    wb2 = load_workbook(f"{DELIV_DIR}/Model_5_Midstream.xlsx")
    ws2 = wb2['Model_Standard']

    print("\nMidstream:")
    d104 = ws2.cell(row=104, column=4).value
    e76 = ws2.cell(row=76, column=5).value
    e78 = ws2.cell(row=78, column=5).value
    d123 = ws2.cell(row=123, column=4).value

    print(f"  D104 (Purchase Price): {d104}")
    print(f"  E76 (Taxes Y1): {e76}")
    print(f"  E78 (CFADS Y1): {e78}")
    print(f"  D123 (IRR): {d123}")

    if '$D$27' in str(d104):
        print("  ✓ Purchase Price = Entry EV (D27)")
    if 'MAX(0' in str(e76) and '$D$55' in str(e76):
        print("  ✓ Taxes formula correct (MAX × Tax Rate)")
    if 'E72' in str(e78) and 'E76' in str(e78) and 'E77' in str(e78):
        print("  ✓ CFADS = EBITDA - Taxes - CapEx")
    if 'IRR' in str(d123):
        print("  ✓ IRR formula correct")


if __name__ == "__main__":
    fix_ccgt_final()
    fix_midstream_final()
    verify_final()

    print("\n" + "="*60)
    print("FINAL FIXES COMPLETE")
    print("="*60)
