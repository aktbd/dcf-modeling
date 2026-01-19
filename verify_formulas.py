#!/usr/bin/env python3
"""
Verify key formulas in the CCGT model have correct cell references
"""

import openpyxl

def verify_ccgt_model(filepath):
    """Verify formulas are correctly structured in CCGT model"""
    print(f"\nVerifying formulas in: {filepath}")
    print("="*60)

    wb = openpyxl.load_workbook(filepath, data_only=False)
    ws = wb.active

    # Build a dictionary of row labels to row numbers
    labels = {}
    for row in range(1, 150):
        label = ws.cell(row=row, column=1).value
        if label:
            labels[str(label).strip()] = row

    print("\nKey row positions found:")
    key_rows = ['Capacity', 'Heat Rate', 'Capacity Factor', 'Forced Outage Rate',
                'Power Price Y1', 'Capacity Price Y1', 'Gas Price Y1', 'Escalation',
                'UCAP', 'Annual Generation', 'Debt Amount', 'EBITDA', 'CFADS',
                'Beginning Balance', 'Interest', 'DSCR', 'Levered IRR']
    for kr in key_rows:
        for label, row in labels.items():
            if kr.lower() in label.lower():
                val_or_formula = ws.cell(row=row, column=4).value
                if val_or_formula and str(val_or_formula).startswith('='):
                    print(f"  {label} (Row {row}): {val_or_formula[:60]}...")
                else:
                    print(f"  {label} (Row {row}): {val_or_formula}")
                break

    # Verify specific critical formulas
    print("\n" + "="*60)
    print("VERIFYING CRITICAL FORMULAS")
    print("="*60)

    errors = []

    # Find specific rows
    ucap_row = None
    gen_row = None
    fuel_row = None
    ebitda_row = None

    for label, row in labels.items():
        if 'UCAP' in label and 'Unforced' in label:
            ucap_row = row
        elif 'Annual Generation' in label:
            gen_row = row
        elif 'Fuel Cost' in label:
            fuel_row = row
        elif label == 'EBITDA':
            ebitda_row = row

    # 1. Verify UCAP formula references Capacity and Forced Outage Rate
    if ucap_row:
        ucap_formula = str(ws.cell(row=ucap_row, column=4).value)
        print(f"\n1. UCAP Formula (Row {ucap_row}):")
        print(f"   {ucap_formula}")

        # Should reference Capacity row and Forced Outage Rate row
        cap_row = labels.get('Capacity')
        for_row = labels.get('Forced Outage Rate')

        if cap_row and f"$D${cap_row}" in ucap_formula:
            print(f"   ✓ Correctly references Capacity at row {cap_row}")
        else:
            errors.append(f"UCAP doesn't reference Capacity correctly")
            print(f"   ✗ Missing reference to Capacity (expected $D${cap_row})")

        if for_row and f"$D${for_row}" in ucap_formula:
            print(f"   ✓ Correctly references Forced Outage Rate at row {for_row}")
        else:
            errors.append(f"UCAP doesn't reference Forced Outage Rate correctly")
            print(f"   ✗ Missing reference to Forced Outage Rate")

    # 2. Verify Generation formula
    if gen_row:
        gen_formula = str(ws.cell(row=gen_row, column=4).value)
        print(f"\n2. Generation Formula (Row {gen_row}):")
        print(f"   {gen_formula}")

        cap_row = labels.get('Capacity')
        cf_row = labels.get('Capacity Factor')

        if '8760' in gen_formula:
            print(f"   ✓ Contains 8760 hours/year")
        else:
            errors.append("Generation doesn't multiply by 8760")

        if cap_row and f"$D${cap_row}" in gen_formula:
            print(f"   ✓ Correctly references Capacity")

        if cf_row and f"$D${cf_row}" in gen_formula:
            print(f"   ✓ Correctly references Capacity Factor")

    # 3. Verify Fuel Cost formula (critical dimensional analysis)
    if fuel_row:
        fuel_formula_e = str(ws.cell(row=fuel_row, column=5).value)  # Year 1
        print(f"\n3. Fuel Cost Formula Y1 (Row {fuel_row}):")
        print(f"   {fuel_formula_e}")

        # Should divide by 1 billion (1000000000 or 1e9)
        if '1000000000' in fuel_formula_e or '1E9' in fuel_formula_e.upper() or '1E+9' in fuel_formula_e.upper():
            print(f"   ✓ Correctly divides by 1 billion")
        else:
            errors.append("Fuel Cost may have dimensional error - missing /1e9")
            print(f"   ✗ WARNING: May have dimensional error (should divide by 1e9)")

        # Should reference Generation, Heat Rate, and Gas Price
        hr_row = labels.get('Heat Rate')
        if hr_row and f"$D${hr_row}" in fuel_formula_e:
            print(f"   ✓ References Heat Rate")

    # 4. Check DSCR formula
    dscr_row = None
    for label, row in labels.items():
        if label == 'DSCR':
            dscr_row = row
            break

    if dscr_row:
        dscr_formula = str(ws.cell(row=dscr_row, column=5).value)
        print(f"\n4. DSCR Formula Y1 (Row {dscr_row}):")
        print(f"   {dscr_formula}")

        if 'IF(' in dscr_formula.upper():
            print(f"   ✓ Has IF statement to prevent #DIV/0")
        else:
            print(f"   ! Consider adding IF to prevent #DIV/0")

    # 5. Check IRR formula
    irr_row = None
    for label, row in labels.items():
        if 'Levered IRR' in label:
            irr_row = row
            break

    if irr_row:
        irr_formula = str(ws.cell(row=irr_row, column=4).value)
        print(f"\n5. Levered IRR Formula (Row {irr_row}):")
        print(f"   {irr_formula}")

        if 'IRR(' in irr_formula.upper():
            print(f"   ✓ Uses IRR function")

    # 6. Check Sources & Uses balance
    su_check_row = None
    for label, row in labels.items():
        if 'Balance Check' in label:
            su_check_row = row
            break

    if su_check_row:
        su_formula = str(ws.cell(row=su_check_row, column=4).value)
        print(f"\n6. S&U Balance Check Formula (Row {su_check_row}):")
        print(f"   {su_formula}")

        if 'ABS(' in su_formula.upper() and 'IF(' in su_formula.upper():
            print(f"   ✓ Uses ABS with tolerance check")

    # Summary
    print("\n" + "="*60)
    print("VERIFICATION SUMMARY")
    print("="*60)

    if errors:
        print(f"\n❌ ERRORS FOUND: {len(errors)}")
        for e in errors:
            print(f"   - {e}")
        return False
    else:
        print("\n✓ All critical formulas verified")
        return True

if __name__ == "__main__":
    filepath = "/mnt/user-data/outputs/Model_1_CCGT.xlsx"
    success = verify_ccgt_model(filepath)

    if success:
        print("\nModel verification PASSED")
    else:
        print("\nModel verification FAILED - review errors above")
