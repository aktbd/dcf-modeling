#!/usr/bin/env python3
"""
Add IC Memo tabs to Model files by reading content from IC_Memo_*.docx files.
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from docx import Document
import os

# File pairs: (Model file, IC Memo file)
BASE_DIR = "/home/user/dcf-modeling/deliverables/deliverables-20260120-0030"

pairs = [
    ("Model_1_CCGT.xlsx", "IC_Memo_1_CCGT.docx"),
    ("Model_2_Peaker.xlsx", "IC_Memo_2_Peaker.docx"),
    ("Model_3_SolarBESS.xlsx", "IC_Memo_3_SolarBESS.docx"),
    ("Model_4_Transmission.xlsx", "IC_Memo_4_Transmission.docx"),
    ("Model_5_Midstream.xlsx", "IC_Memo_5_Midstream.docx"),
]

def extract_docx_content(docx_path):
    """Extract all text content from a Word document."""
    doc = Document(docx_path)
    content = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Detect headings by style
            style_name = para.style.name if para.style else ""
            is_heading = "Heading" in style_name or para.runs and any(run.bold for run in para.runs)
            content.append((text, is_heading, style_name))

    # Also extract tables
    for table in doc.tables:
        content.append(("", False, ""))  # Blank line before table
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            content.append((row_text, False, "table"))

    return content

def add_ic_memo_sheet(xlsx_path, docx_path):
    """Add IC Memo sheet to Excel workbook with content from Word doc."""
    print(f"Processing: {os.path.basename(xlsx_path)}")

    # Load workbook
    wb = load_workbook(xlsx_path)

    # Remove existing IC Memo sheet if present
    if "IC Memo" in wb.sheetnames:
        del wb["IC Memo"]

    # Create new IC Memo sheet
    ws = wb.create_sheet("IC Memo")

    # Extract content from Word doc
    content = extract_docx_content(docx_path)

    # Style definitions
    header_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    subheader_font = Font(name="Calibri", size=12, bold=True)
    body_font = Font(name="Calibri", size=11)
    table_font = Font(name="Consolas", size=10)

    # Write content to sheet
    row = 1
    for text, is_heading, style_name in content:
        cell = ws.cell(row=row, column=1, value=text)

        if "Heading 1" in style_name or (is_heading and len(text) < 50):
            cell.font = header_font
        elif "Heading" in style_name or is_heading:
            cell.font = subheader_font
        elif style_name == "table":
            cell.font = table_font
        else:
            cell.font = body_font

        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    # Set column width
    ws.column_dimensions["A"].width = 120

    # Save workbook
    wb.save(xlsx_path)
    print(f"  Added IC Memo sheet with {row-1} rows")

def main():
    print("=" * 60)
    print("Adding IC Memo tabs to Model files")
    print("=" * 60)

    for model_file, memo_file in pairs:
        xlsx_path = os.path.join(BASE_DIR, model_file)
        docx_path = os.path.join(BASE_DIR, memo_file)

        if os.path.exists(xlsx_path) and os.path.exists(docx_path):
            add_ic_memo_sheet(xlsx_path, docx_path)
        else:
            print(f"SKIP: Missing files for {model_file}")

    print("=" * 60)
    print("Verifying sheets in each Model file:")
    print("=" * 60)

    for model_file, _ in pairs:
        xlsx_path = os.path.join(BASE_DIR, model_file)
        wb = load_workbook(xlsx_path, read_only=True)
        print(f"  {model_file}: {wb.sheetnames}")
        wb.close()

    print("\nDone!")

if __name__ == "__main__":
    main()
