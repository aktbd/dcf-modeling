#!/usr/bin/env python3
"""
Verify audit fixes and regenerate ZIP
"""

from openpyxl import load_workbook
import os
import zipfile
import subprocess
import tempfile

os.chdir("/home/user/dcf-modeling/deliverables/deliverables-20260120-0030")

print("=" * 70)
print("AUDIT FIX VERIFICATION")
print("=" * 70)

# Verify Midstream Revenue formula (should NOT have *1000)
print("\n1. MIDSTREAM REVENUE CHECK")
wb = load_workbook("Model_5_Midstream.xlsx")
ws = wb["Model_Standard"]
# Find Revenue row (typically around row 66-70 after control panel inserted)
revenue_found = False
for row in range(50, 150):
    label = ws.cell(row=row, column=1).value
    if label and "Total Revenue" in str(label):
        formula = ws.cell(row=row, column=5).value  # E column
        print(f"   Row {row}: {label}")
        print(f"   Formula E{row}: {formula}")
        if formula and "*1000" in str(formula):
            print("   ❌ FAIL: Still has *1000 error!")
        else:
            print("   ✓ PASS: No *1000 in formula")
        revenue_found = True
        break
if not revenue_found:
    # Check for Gathering Revenue
    for row in range(50, 150):
        label = ws.cell(row=row, column=1).value
        if label and ("Gathering" in str(label) or "Revenue" in str(label)):
            formula = ws.cell(row=row, column=5).value
            print(f"   Row {row}: {label}")
            print(f"   Formula E{row}: {formula}")
wb.close()

# Verify Transmission S&U (DSRA should not be in Sources)
print("\n2. TRANSMISSION S&U CHECK")
wb = load_workbook("Model_4_Transmission.xlsx")
ws = wb["Model_Standard"]
# Find Total Sources row
for row in range(100, 150):
    label = ws.cell(row=row, column=1).value
    if label and "Total Sources" in str(label):
        formula = ws.cell(row=row, column=4).value  # D column
        print(f"   Row {row}: {label}")
        print(f"   Formula D{row}: {formula}")
        if formula and ("D104" in str(formula) or "D114" in str(formula)):
            print("   ❌ FAIL: DSRA still in Total Sources!")
        else:
            print("   ✓ PASS: DSRA not in Total Sources")
        break
wb.close()

# Verify Solar Tax formula (should have MAX(0,...))
print("\n3. SOLAR TAX CHECK")
wb = load_workbook("Model_3_SolarBESS.xlsx")
ws = wb["Model_Standard"]
# Find Tax Expense row
for row in range(60, 120):
    label = ws.cell(row=row, column=1).value
    if label and "Tax Expense" in str(label):
        formula = ws.cell(row=row, column=5).value  # E column
        print(f"   Row {row}: {label}")
        print(f"   Formula E{row}: {formula}")
        if formula and "MAX(0" in str(formula):
            print("   ✓ PASS: Has MAX(0,...) guardrail")
        elif formula:
            print("   ⚠ WARN: No MAX(0) in formula")
        break
wb.close()

# Verify Control Panel exists
print("\n4. CONTROL PANEL CHECK")
wb = load_workbook("Model_1_CCGT.xlsx")
ws = wb["Model_Standard"]
control_found = False
for row in range(1, 20):
    val = ws.cell(row=row, column=1).value
    if val and "CONTROL PANEL" in str(val):
        print(f"   ✓ Found Control Panel at row {row}")
        control_found = True
        # Print flags
        for r in range(row+1, row+10):
            label = ws.cell(row=r, column=1).value
            value = ws.cell(row=r, column=4).value
            if label and value:
                print(f"      {label}: {value}")
        break
if not control_found:
    print("   ❌ Control Panel not found!")
wb.close()

# Verify LLCR/PLCR in Model_Full
print("\n5. LLCR/PLCR CHECK")
wb = load_workbook("Model_1_CCGT.xlsx")
ws = wb["Model_Full"]
llcr_found = False
for row in range(100, 150):
    val = ws.cell(row=row, column=1).value
    if val and "LLCR" in str(val):
        formula = ws.cell(row=row, column=4).value
        print(f"   ✓ Found LLCR at row {row}")
        print(f"     Formula: {formula}")
        llcr_found = True
        break
if not llcr_found:
    print("   ⚠ LLCR not found in expected location")
wb.close()

# Verify Model_Quick exists and has key metrics
print("\n6. MODEL_QUICK CHECK")
wb = load_workbook("Model_1_CCGT.xlsx")
ws = wb["Model_Quick"]
print(f"   ✓ Model_Quick exists")
print(f"   Max row: {ws.max_row}")
# Find IRR
for row in range(1, ws.max_row + 1):
    val = ws.cell(row=row, column=1).value
    if val and "Levered IRR" in str(val):
        formula = ws.cell(row=row, column=4).value
        print(f"   ✓ Levered IRR at row {row}: {formula}")
        break
wb.close()

print("\n" + "=" * 70)
print("REGENERATING ZIP")
print("=" * 70)

# Delete old ZIP
zip_path = "deliverables_bundle.zip"
if os.path.exists(zip_path):
    os.remove(zip_path)
    print(f"  Deleted old {zip_path}")

# Create new ZIP
files_to_zip = [
    "Model_1_CCGT.xlsx", "Model_2_Peaker.xlsx", "Model_3_SolarBESS.xlsx",
    "Model_4_Transmission.xlsx", "Model_5_Midstream.xlsx",
    "Drill_1_CCGT.xlsx", "Drill_2_Peaker.xlsx", "Drill_3_SolarBESS.xlsx",
    "Drill_4_Transmission.xlsx", "Drill_5_Midstream.xlsx",
    "IC_Memo_1_CCGT.docx", "IC_Memo_2_Peaker.docx", "IC_Memo_3_SolarBESS.docx",
    "IC_Memo_4_Transmission.docx", "IC_Memo_5_Midstream.docx",
    "Master_Template.xlsx", "Lotus_Interview_CheatSheet_Expanded.pdf",
    "Lotus_Interview_CheatSheet_Expanded.md", "Implementation_Manifesto.md",
]

with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
    for fname in files_to_zip:
        if os.path.exists(fname):
            zf.write(fname)
            print(f"  Added: {fname}")
        else:
            print(f"  SKIP: {fname} not found")

print(f"\n  Created: {zip_path}")

# Verify by extraction
print("\n--- VERIFICATION BY EXTRACTION ---")
with tempfile.TemporaryDirectory() as tmpdir:
    with zipfile.ZipFile(zip_path, 'r') as zf:
        zf.extractall(tmpdir)

    # Check a model file
    test_file = os.path.join(tmpdir, "Model_1_CCGT.xlsx")
    wb = load_workbook(test_file, read_only=True)
    print(f"  Model_1_CCGT.xlsx sheets: {wb.sheetnames}")
    wb.close()

    # Check Midstream
    test_file = os.path.join(tmpdir, "Model_5_Midstream.xlsx")
    wb = load_workbook(test_file, read_only=True)
    print(f"  Model_5_Midstream.xlsx sheets: {wb.sheetnames}")
    wb.close()

    # Check a drill
    test_file = os.path.join(tmpdir, "Drill_1_CCGT.xlsx")
    wb = load_workbook(test_file, read_only=True)
    print(f"  Drill_1_CCGT.xlsx sheets: {wb.sheetnames}")
    wb.close()

print("\n  ✓ ZIP verified by extraction!")

# SHA256
result = subprocess.run(['sha256sum', zip_path], capture_output=True, text=True)
print(f"\n  SHA256: {result.stdout.strip()}")

print("\n" + "=" * 70)
print("COMPLETE")
print("=" * 70)
