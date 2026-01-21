#!/usr/bin/env python3
"""
Enterprise Gas Peaker Model - SDG&E Tolling Agreement
Single-tab, formula-only, annual model (2014-2033)
Pre-tax equity returns for Life Hold and Short Hold scenarios
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# Style definitions
BLUE_FONT = Font(color="0000FF", bold=True)
BLACK_FONT = Font(color="000000")
GREEN_FONT = Font(color="006400")
BOLD_FONT = Font(bold=True)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
INPUT_FILL = PatternFill(start_color="FFFFC7", end_color="FFFFC7", fill_type="solid")
OUTPUT_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
THIN_BORDER = Border(bottom=Side(style='thin'))
DOUBLE_BORDER = Border(bottom=Side(style='double'))

# Number formats
CURRENCY_FMT = '_($* #,##0_);_($* (#,##0);_($* "-"_);_(@_)'
CURRENCY_DEC_FMT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"_);_(@_)'
PERCENT_FMT = '0.00%'
MULTIPLE_FMT = '0.00"x"'
NUMBER_FMT = '#,##0'
DECIMAL_FMT = '#,##0.00'


def apply_input_style(ws, row, col, value=None):
    cell = ws.cell(row=row, column=col)
    cell.font = BLUE_FONT
    cell.fill = INPUT_FILL
    if value is not None:
        cell.value = value
    return cell


def apply_calc_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    cell.font = BLACK_FONT
    if formula is not None:
        cell.value = formula
    return cell


def apply_output_style(ws, row, col, formula=None):
    cell = ws.cell(row=row, column=col)
    cell.font = Font(color="000000", bold=True)
    cell.fill = OUTPUT_FILL
    if formula is not None:
        cell.value = formula
    return cell


def set_label(ws, row, label, col=1):
    cell = ws.cell(row=row, column=col)
    cell.value = label
    return cell


def set_header(ws, row, text, col=1, span=1):
    cell = ws.cell(row=row, column=col)
    cell.value = text
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    return cell


def col_letter(col):
    return get_column_letter(col)


def build_model():
    wb = Workbook()
    ws = wb.active
    ws.title = "Enterprise Peaker"

    # Column layout: A=Labels, B=Units, C=Notes, D=Y0(2013), E=Y1(2014), ... X=Y20(2033)
    # Years: 2013 (Y0), 2014-2033 (Y1-Y20)

    rows = {}
    row = 1

    # Set column widths
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 20
    for col in range(4, 25):  # D through X
        ws.column_dimensions[col_letter(col)].width = 14

    # Freeze panes
    ws.freeze_panes = 'E3'

    # ============================================================
    # SECTION 1: INPUTS
    # ============================================================
    set_header(ws, row, "═══ INPUTS ═══")
    row += 1

    # Year headers
    ws.cell(row=row, column=1).value = "Year"
    ws.cell(row=row, column=4).value = 2013
    for yr in range(2014, 2034):
        ws.cell(row=row, column=5 + yr - 2014).value = yr
    rows['year_header'] = row
    row += 1

    # Timeline inputs
    set_header(ws, row, "Timeline")
    row += 1

    rows['year_index'] = row
    set_label(ws, row, "Year Index")
    ws.cell(row=row, column=2).value = "yr"
    ws.cell(row=row, column=4).value = 0
    for col in range(5, 25):
        ws.cell(row=row, column=col).value = col - 4
    row += 1

    rows['close_year'] = row
    set_label(ws, row, "Close Date")
    apply_input_style(ws, row, 4, "12/31/2013")
    row += 2

    # Plant / Dispatch inputs
    set_header(ws, row, "Plant / Dispatch")
    row += 1

    rows['capacity_mw'] = row
    set_label(ws, row, "Contract Capacity (net)")
    ws.cell(row=row, column=2).value = "MW"
    apply_input_style(ws, row, 4, 51)
    row += 1

    rows['capacity_kw'] = row
    set_label(ws, row, "Contract Capacity")
    ws.cell(row=row, column=2).value = "kW"
    apply_calc_style(ws, row, 4, "=D" + str(rows['capacity_mw']) + "*1000")
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['billing_capacity'] = row
    set_label(ws, row, "Billing Capacity")
    ws.cell(row=row, column=2).value = "kW"
    apply_calc_style(ws, row, 4, "=ROUND(D" + str(rows['capacity_kw']) + ",-2)")
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['capacity_factor'] = row
    set_label(ws, row, "Capacity Factor")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.05)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['starts_per_month'] = row
    set_label(ws, row, "Starts per Month")
    ws.cell(row=row, column=2).value = "#"
    apply_input_style(ws, row, 4, 5)
    row += 1

    rows['annual_starts'] = row
    set_label(ws, row, "Annual Starts")
    ws.cell(row=row, column=2).value = "#"
    apply_calc_style(ws, row, 4, "=D" + str(rows['starts_per_month']) + "*12")
    row += 1

    rows['eaf'] = row
    set_label(ws, row, "EAF (Availability)")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.985)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['cpi'] = row
    set_label(ws, row, "CPI Escalation")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.025)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['libor'] = row
    set_label(ws, row, "LIBOR (Base)")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.03)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 2

    # Contract Revenue inputs
    set_header(ws, row, "Contract Revenue Rates")
    row += 1

    rows['capacity_rate'] = row
    set_label(ws, row, "Capacity Rate")
    ws.cell(row=row, column=2).value = "$/kW-mo"
    apply_input_style(ws, row, 4, 9.17)
    ws.cell(row=row, column=4).number_format = DECIMAL_FMT
    row += 1

    rows['vom_rate'] = row
    set_label(ws, row, "VOM Payment Rate")
    ws.cell(row=row, column=2).value = "$/MWh"
    apply_input_style(ws, row, 4, 5.00)
    ws.cell(row=row, column=4).number_format = DECIMAL_FMT
    row += 1

    rows['start_payment'] = row
    set_label(ws, row, "Start-Up Payment")
    ws.cell(row=row, column=2).value = "$/start"
    apply_input_style(ws, row, 4, 750)
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['hrp_toggle'] = row
    set_label(ws, row, "HRP Toggle (0=Off, 1=On)")
    ws.cell(row=row, column=2).value = "flag"
    apply_input_style(ws, row, 4, 0)
    row += 2

    # Opex inputs
    set_header(ws, row, "Operating Cost Rates")
    row += 1

    rows['fom_rate'] = row
    set_label(ws, row, "FOM Rate")
    ws.cell(row=row, column=2).value = "$/kW-yr"
    apply_input_style(ws, row, 4, 25.00)
    ws.cell(row=row, column=4).number_format = DECIMAL_FMT
    row += 1

    rows['internal_vom_rate'] = row
    set_label(ws, row, "Internal VOM Expense")
    ws.cell(row=row, column=2).value = "$/MWh"
    apply_input_style(ws, row, 4, 3.00)
    ws.cell(row=row, column=4).number_format = DECIMAL_FMT
    row += 1

    rows['internal_start_cost'] = row
    set_label(ws, row, "Internal Start Cost")
    ws.cell(row=row, column=2).value = "$/start"
    apply_input_style(ws, row, 4, 300)
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['lc_notional'] = row
    set_label(ws, row, "LC Notional")
    ws.cell(row=row, column=2).value = "$"
    apply_input_style(ws, row, 4, 2900000)
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    row += 1

    rows['lc_fee'] = row
    set_label(ws, row, "LC Fee")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.03)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['lc_cost'] = row
    set_label(ws, row, "Annual LC Cost")
    ws.cell(row=row, column=2).value = "$"
    apply_calc_style(ws, row, 4, "=D" + str(rows['lc_notional']) + "*D" + str(rows['lc_fee']))
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    row += 2

    # Acquisition / Exit inputs
    set_header(ws, row, "Acquisition / Exit")
    row += 1

    rows['purchase_price_kw'] = row
    set_label(ws, row, "Purchase Price")
    ws.cell(row=row, column=2).value = "$/kW"
    apply_input_style(ws, row, 4, 850)
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['purchase_price'] = row
    set_label(ws, row, "Total Purchase Price")
    ws.cell(row=row, column=2).value = "$"
    apply_calc_style(ws, row, 4, "=D" + str(rows['purchase_price_kw']) + "*D" + str(rows['capacity_kw']))
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    row += 1

    rows['residual_kw'] = row
    set_label(ws, row, "Residual Value")
    ws.cell(row=row, column=2).value = "$/kW"
    apply_input_style(ws, row, 4, 250)
    ws.cell(row=row, column=4).number_format = NUMBER_FMT
    row += 1

    rows['residual_value'] = row
    set_label(ws, row, "Total Residual Value")
    ws.cell(row=row, column=2).value = "$"
    apply_calc_style(ws, row, 4, "=D" + str(rows['residual_kw']) + "*D" + str(rows['capacity_kw']))
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    row += 1

    rows['buyer_coe'] = row
    set_label(ws, row, "Buyer CoE (Short Hold)")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.10)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['short_hold_exit_year'] = row
    set_label(ws, row, "Short Hold Exit Year")
    ws.cell(row=row, column=2).value = "yr"
    apply_input_style(ws, row, 4, 5)
    row += 2

    # Debt inputs
    set_header(ws, row, "Debt - Tranche 2 (Subordinate)")
    row += 1

    rows['tr2_principal'] = row
    set_label(ws, row, "Tr2 Principal")
    ws.cell(row=row, column=2).value = "$"
    apply_input_style(ws, row, 4, 2500000)
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    row += 1

    rows['tr2_spread'] = row
    set_label(ws, row, "Tr2 Spread over LIBOR")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.055)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['tr2_rate'] = row
    set_label(ws, row, "Tr2 Interest Rate")
    ws.cell(row=row, column=2).value = "%"
    apply_calc_style(ws, row, 4, "=$D$" + str(rows['libor']) + "+$D$" + str(rows['tr2_spread']))
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['tr2_amort'] = row
    set_label(ws, row, "Tr2 Mandatory Amort")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.01)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 2

    set_header(ws, row, "Debt - Tranche 1 (Senior)")
    row += 1

    rows['tr1_spread'] = row
    set_label(ws, row, "Tr1 Spread over LIBOR")
    ws.cell(row=row, column=2).value = "%"
    apply_input_style(ws, row, 4, 0.03)
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['tr1_rate'] = row
    set_label(ws, row, "Tr1 Interest Rate")
    ws.cell(row=row, column=2).value = "%"
    apply_calc_style(ws, row, 4, "=$D$" + str(rows['libor']) + "+$D$" + str(rows['tr1_spread']))
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['target_dscr'] = row
    set_label(ws, row, "Target DSCR")
    ws.cell(row=row, column=2).value = "x"
    apply_input_style(ws, row, 4, 1.40)
    ws.cell(row=row, column=4).number_format = MULTIPLE_FMT
    row += 1

    rows['tr1_tenor'] = row
    set_label(ws, row, "Tr1 Tenor")
    ws.cell(row=row, column=2).value = "years"
    apply_input_style(ws, row, 4, 20)
    row += 2

    # ============================================================
    # SECTION 2: VOLUMES
    # ============================================================
    set_header(ws, row, "═══ VOLUMES ═══")
    row += 1

    rows['net_mwh'] = row
    set_label(ws, row, "Net MWh (annual)")
    ws.cell(row=row, column=2).value = "MWh"
    for col in range(5, 25):
        apply_calc_style(ws, row, col, "=$D$" + str(rows['capacity_mw']) + "*8760*$D$" + str(rows['capacity_factor']))
        ws.cell(row=row, column=col).number_format = NUMBER_FMT
    row += 1

    rows['starts'] = row
    set_label(ws, row, "Annual Starts")
    ws.cell(row=row, column=2).value = "#"
    for col in range(5, 25):
        apply_calc_style(ws, row, col, "=$D$" + str(rows['annual_starts']))
        ws.cell(row=row, column=col).number_format = NUMBER_FMT
    row += 2

    # ============================================================
    # SECTION 3: AVAILABILITY FACTOR
    # ============================================================
    set_header(ws, row, "═══ AVAILABILITY ADJUSTMENT ═══")
    row += 1

    rows['aaf'] = row
    set_label(ws, row, "AAF (Adjusted Availability)")
    ws.cell(row=row, column=2).value = "%"
    ws.cell(row=row, column=3).value = "EAF-based formula"
    # AAF formula: IF EAF<=0.98: EAF/0.98, IF 0.98<EAF<0.99: 1, IF EAF>=0.99: EAF/0.99
    for col in range(5, 25):
        formula = f'=IF($D${rows["eaf"]}<=0.98,$D${rows["eaf"]}/0.98,IF($D${rows["eaf"]}<0.99,1,$D${rows["eaf"]}/0.99))'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = PERCENT_FMT
    row += 2

    # ============================================================
    # SECTION 4: REVENUES
    # ============================================================
    set_header(ws, row, "═══ REVENUES ═══")
    row += 1

    rows['capacity_revenue'] = row
    set_label(ws, row, "Capacity Revenue")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Flat rate × AAF"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'=$D${rows["billing_capacity"]}*$D${rows["capacity_rate"]}*12*{c}{rows["aaf"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['vom_revenue'] = row
    set_label(ws, row, "VOM Revenue")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Escalates by CPI"
    for col in range(5, 25):
        c = col_letter(col)
        yr_idx = col - 4
        formula = f'={c}{rows["net_mwh"]}*$D${rows["vom_rate"]}*(1+$D${rows["cpi"]})^({yr_idx}-1)'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['start_revenue'] = row
    set_label(ws, row, "Start-Up Revenue")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Escalates by CPI"
    for col in range(5, 25):
        c = col_letter(col)
        yr_idx = col - 4
        formula = f'={c}{rows["starts"]}*$D${rows["start_payment"]}*(1+$D${rows["cpi"]})^({yr_idx}-1)'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['hrp'] = row
    set_label(ws, row, "Heat Rate Payment (HRP)")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Base case = 0"
    for col in range(5, 25):
        formula = f'=$D${rows["hrp_toggle"]}*0'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['total_revenue'] = row
    set_label(ws, row, "Total Revenue")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=1).font = BOLD_FONT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["capacity_revenue"]}+{c}{rows["vom_revenue"]}+{c}{rows["start_revenue"]}+{c}{rows["hrp"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=row, column=col).font = BOLD_FONT
        ws.cell(row=row, column=col).border = THIN_BORDER
    row += 2

    # ============================================================
    # SECTION 5: OPEX
    # ============================================================
    set_header(ws, row, "═══ OPERATING EXPENSES ═══")
    row += 1

    rows['fom'] = row
    set_label(ws, row, "FOM Expense")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Escalates by CPI"
    for col in range(5, 25):
        yr_idx = col - 4
        formula = f'=$D${rows["capacity_kw"]}*$D${rows["fom_rate"]}*(1+$D${rows["cpi"]})^({yr_idx}-1)'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['internal_vom'] = row
    set_label(ws, row, "Internal VOM Expense")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Escalates by CPI"
    for col in range(5, 25):
        c = col_letter(col)
        yr_idx = col - 4
        formula = f'={c}{rows["net_mwh"]}*$D${rows["internal_vom_rate"]}*(1+$D${rows["cpi"]})^({yr_idx}-1)'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['internal_starts'] = row
    set_label(ws, row, "Internal Start Expense")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Escalates by CPI"
    for col in range(5, 25):
        c = col_letter(col)
        yr_idx = col - 4
        formula = f'={c}{rows["starts"]}*$D${rows["internal_start_cost"]}*(1+$D${rows["cpi"]})^({yr_idx}-1)'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['lc_expense'] = row
    set_label(ws, row, "LC Cost")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Flat"
    for col in range(5, 25):
        formula = f'=$D${rows["lc_cost"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['total_opex'] = row
    set_label(ws, row, "Total Opex")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=1).font = BOLD_FONT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["fom"]}+{c}{rows["internal_vom"]}+{c}{rows["internal_starts"]}+{c}{rows["lc_expense"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=row, column=col).font = BOLD_FONT
        ws.cell(row=row, column=col).border = THIN_BORDER
    row += 2

    # ============================================================
    # SECTION 6: EBITDA & CFADS
    # ============================================================
    set_header(ws, row, "═══ EBITDA & CFADS ═══")
    row += 1

    rows['ebitda'] = row
    set_label(ws, row, "EBITDA")
    ws.cell(row=row, column=2).value = "$"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["total_revenue"]}-{c}{rows["total_opex"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=row, column=col).font = BOLD_FONT
    row += 1

    rows['wc'] = row
    set_label(ws, row, "Working Capital Change")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Assume = 0"
    for col in range(5, 25):
        apply_calc_style(ws, row, col, 0)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['capex'] = row
    set_label(ws, row, "CapEx")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Assume = 0"
    for col in range(5, 25):
        apply_calc_style(ws, row, col, 0)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['cfads'] = row
    set_label(ws, row, "CFADS")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=1).font = BOLD_FONT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["ebitda"]}-{c}{rows["wc"]}-{c}{rows["capex"]}'
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=row, column=col).border = DOUBLE_BORDER
    row += 2

    # ============================================================
    # SECTION 7: DEBT SCHEDULE - TRANCHE 2
    # ============================================================
    set_header(ws, row, "═══ DEBT SCHEDULE - TRANCHE 2 (SUB) ═══")
    row += 1

    # Define row numbers first for forward references
    tr2_beg_row = row
    tr2_interest_row = row + 1
    tr2_principal_row = row + 2
    tr2_ds_row = row + 3
    tr2_end_row = row + 4

    rows['tr2_beg'] = tr2_beg_row
    set_label(ws, tr2_beg_row, "Tr2 Beginning Balance")
    ws.cell(row=tr2_beg_row, column=2).value = "$"
    apply_calc_style(ws, tr2_beg_row, 4, f'=$D${rows["tr2_principal"]}')
    ws.cell(row=tr2_beg_row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 25):
        prev_col = col_letter(col - 1)
        formula = f'={prev_col}{tr2_end_row}'
        apply_calc_style(ws, tr2_beg_row, col, formula)
        ws.cell(row=tr2_beg_row, column=col).number_format = CURRENCY_FMT

    rows['tr2_interest'] = tr2_interest_row
    set_label(ws, tr2_interest_row, "Tr2 Interest")
    ws.cell(row=tr2_interest_row, column=2).value = "$"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{tr2_beg_row}*$D${rows["tr2_rate"]}'
        apply_calc_style(ws, tr2_interest_row, col, formula)
        ws.cell(row=tr2_interest_row, column=col).number_format = CURRENCY_FMT

    rows['tr2_principal_pmt'] = tr2_principal_row
    set_label(ws, tr2_principal_row, "Tr2 Principal")
    ws.cell(row=tr2_principal_row, column=2).value = "$"
    ws.cell(row=tr2_principal_row, column=3).value = "1%/yr; balloon final"
    for col in range(5, 24):  # Years 1-19
        c = col_letter(col)
        formula = f'=MIN({c}{tr2_beg_row},$D${rows["tr2_principal"]}*$D${rows["tr2_amort"]})'
        apply_calc_style(ws, tr2_principal_row, col, formula)
        ws.cell(row=tr2_principal_row, column=col).number_format = CURRENCY_FMT
    # Year 20 - balloon payment
    apply_calc_style(ws, tr2_principal_row, 24, f'=X{tr2_beg_row}')
    ws.cell(row=tr2_principal_row, column=24).number_format = CURRENCY_FMT

    rows['tr2_ds'] = tr2_ds_row
    set_label(ws, tr2_ds_row, "Tr2 Debt Service")
    ws.cell(row=tr2_ds_row, column=2).value = "$"
    ws.cell(row=tr2_ds_row, column=1).font = BOLD_FONT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{tr2_interest_row}+{c}{tr2_principal_row}'
        apply_calc_style(ws, tr2_ds_row, col, formula)
        ws.cell(row=tr2_ds_row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=tr2_ds_row, column=col).font = BOLD_FONT
        ws.cell(row=tr2_ds_row, column=col).border = THIN_BORDER

    rows['tr2_end'] = tr2_end_row
    set_label(ws, tr2_end_row, "Tr2 Ending Balance")
    ws.cell(row=tr2_end_row, column=2).value = "$"
    # Y0 ending balance = Y0 beginning balance (no debt service in Y0)
    apply_calc_style(ws, tr2_end_row, 4, f'=D{tr2_beg_row}')
    ws.cell(row=tr2_end_row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'=MAX(0,{c}{tr2_beg_row}-{c}{tr2_principal_row})'
        apply_calc_style(ws, tr2_end_row, col, formula)
        ws.cell(row=tr2_end_row, column=col).number_format = CURRENCY_FMT

    row = tr2_end_row + 2

    # ============================================================
    # SECTION 8: DEBT SIZING - TRANCHE 1 (Senior)
    # ============================================================
    set_header(ws, row, "═══ DEBT SIZING - TRANCHE 1 (SENIOR) ═══")
    row += 1

    rows['max_total_ds'] = row
    set_label(ws, row, "Max Total DS Allowed")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "CFADS / 1.40"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["cfads"]}/$D${rows["target_dscr"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['max_senior_ds'] = row
    set_label(ws, row, "Max Senior DS Allowed")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Max Total - Tr2"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'=MAX(0,{c}{rows["max_total_ds"]}-{c}{rows["tr2_ds"]})'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['tr1_annual_ds'] = row
    set_label(ws, row, "Senior Annual DS (Level)")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "MIN of all years"
    # Constant DS = MIN of max allowed across all years
    apply_calc_style(ws, row, 4, f'=MIN(E{rows["max_senior_ds"]}:X{rows["max_senior_ds"]})')
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    ws.cell(row=row, column=4).fill = OUTPUT_FILL
    ws.cell(row=row, column=4).font = Font(bold=True)
    row += 1

    rows['tr1_debt0'] = row
    set_label(ws, row, "Senior Debt (Sized)")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "PV of level DS"
    # PV of level DS over 20 years at senior rate
    apply_calc_style(ws, row, 4, f'=PV($D${rows["tr1_rate"]},$D${rows["tr1_tenor"]},-$D${rows["tr1_annual_ds"]},0,0)')
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    ws.cell(row=row, column=4).fill = OUTPUT_FILL
    ws.cell(row=row, column=4).font = Font(bold=True)
    row += 2

    # Senior amortization schedule
    set_header(ws, row, "═══ DEBT SCHEDULE - TRANCHE 1 (SENIOR) ═══")
    row += 1

    # Define row numbers first for forward references
    tr1_beg_row = row
    tr1_interest_row = row + 1
    tr1_principal_row = row + 2
    tr1_ds_row = row + 3
    tr1_end_row = row + 4

    rows['tr1_beg'] = tr1_beg_row
    set_label(ws, tr1_beg_row, "Tr1 Beginning Balance")
    ws.cell(row=tr1_beg_row, column=2).value = "$"
    apply_calc_style(ws, tr1_beg_row, 4, f'=$D${rows["tr1_debt0"]}')
    ws.cell(row=tr1_beg_row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 25):
        prev_col = col_letter(col - 1)
        formula = f'={prev_col}{tr1_end_row}'
        apply_calc_style(ws, tr1_beg_row, col, formula)
        ws.cell(row=tr1_beg_row, column=col).number_format = CURRENCY_FMT

    rows['tr1_interest'] = tr1_interest_row
    set_label(ws, tr1_interest_row, "Tr1 Interest")
    ws.cell(row=tr1_interest_row, column=2).value = "$"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{tr1_beg_row}*$D${rows["tr1_rate"]}'
        apply_calc_style(ws, tr1_interest_row, col, formula)
        ws.cell(row=tr1_interest_row, column=col).number_format = CURRENCY_FMT

    rows['tr1_principal_pmt'] = tr1_principal_row
    set_label(ws, tr1_principal_row, "Tr1 Principal")
    ws.cell(row=tr1_principal_row, column=2).value = "$"
    for col in range(5, 24):  # Years 1-19
        c = col_letter(col)
        formula = f'=MIN({c}{tr1_beg_row},$D${rows["tr1_annual_ds"]}-{c}{tr1_interest_row})'
        apply_calc_style(ws, tr1_principal_row, col, formula)
        ws.cell(row=tr1_principal_row, column=col).number_format = CURRENCY_FMT
    # Year 20 - force to zero
    apply_calc_style(ws, tr1_principal_row, 24, f'=X{tr1_beg_row}')
    ws.cell(row=tr1_principal_row, column=24).number_format = CURRENCY_FMT

    rows['tr1_ds'] = tr1_ds_row
    set_label(ws, tr1_ds_row, "Tr1 Debt Service")
    ws.cell(row=tr1_ds_row, column=2).value = "$"
    ws.cell(row=tr1_ds_row, column=1).font = BOLD_FONT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{tr1_interest_row}+{c}{tr1_principal_row}'
        apply_calc_style(ws, tr1_ds_row, col, formula)
        ws.cell(row=tr1_ds_row, column=col).number_format = CURRENCY_FMT
        ws.cell(row=tr1_ds_row, column=col).font = BOLD_FONT
        ws.cell(row=tr1_ds_row, column=col).border = THIN_BORDER

    rows['tr1_end'] = tr1_end_row
    set_label(ws, tr1_end_row, "Tr1 Ending Balance")
    ws.cell(row=tr1_end_row, column=2).value = "$"
    # Y0 ending balance = Y0 beginning balance (no debt service in Y0)
    apply_calc_style(ws, tr1_end_row, 4, f'=D{tr1_beg_row}')
    ws.cell(row=tr1_end_row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'=MAX(0,{c}{tr1_beg_row}-{c}{tr1_principal_row})'
        apply_calc_style(ws, tr1_end_row, col, formula)
        ws.cell(row=tr1_end_row, column=col).number_format = CURRENCY_FMT

    row = tr1_end_row + 2

    # ============================================================
    # SECTION 9: DSCR CHECK
    # ============================================================
    set_header(ws, row, "═══ DSCR CHECK ═══")
    row += 1

    rows['total_ds'] = row
    set_label(ws, row, "Total Debt Service")
    ws.cell(row=row, column=2).value = "$"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["tr1_ds"]}+{c}{rows["tr2_ds"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 1

    rows['dscr'] = row
    set_label(ws, row, "DSCR")
    ws.cell(row=row, column=2).value = "x"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'=IF({c}{rows["total_ds"]}>0,{c}{rows["cfads"]}/{c}{rows["total_ds"]},0)'
        apply_output_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = MULTIPLE_FMT
    row += 1

    rows['min_dscr'] = row
    set_label(ws, row, "Min DSCR")
    ws.cell(row=row, column=2).value = "x"
    apply_output_style(ws, row, 4, f'=MIN(E{rows["dscr"]}:X{rows["dscr"]})')
    ws.cell(row=row, column=4).number_format = MULTIPLE_FMT
    row += 2

    # ============================================================
    # SECTION 10: EQUITY CASH FLOWS
    # ============================================================
    set_header(ws, row, "═══ EQUITY CASH FLOWS ═══")
    row += 1

    rows['initial_equity'] = row
    set_label(ws, row, "Initial Equity")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "Purchase - Debt"
    apply_calc_style(ws, row, 4, f'=$D${rows["purchase_price"]}-$D${rows["tr1_debt0"]}-$D${rows["tr2_principal"]}')
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    ws.cell(row=row, column=4).fill = OUTPUT_FILL
    row += 1

    rows['annual_equity_cf'] = row
    set_label(ws, row, "Annual Equity CF")
    ws.cell(row=row, column=2).value = "$"
    for col in range(5, 25):
        c = col_letter(col)
        formula = f'={c}{rows["cfads"]}-{c}{rows["total_ds"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    row += 2

    # ============================================================
    # SECTION 11: EXIT VALUES
    # ============================================================
    set_header(ws, row, "═══ EXIT VALUES ═══")
    row += 1

    # Life Hold
    rows['life_hold_residual'] = row
    set_label(ws, row, "Life Hold: Residual (Y20)")
    ws.cell(row=row, column=2).value = "$"
    apply_calc_style(ws, row, 24, f'=$D${rows["residual_value"]}')
    ws.cell(row=row, column=24).number_format = CURRENCY_FMT
    ws.cell(row=row, column=24).fill = OUTPUT_FILL
    row += 1

    # Life Hold Equity CF
    rows['life_hold_equity_cf'] = row
    set_label(ws, row, "Life Hold: Equity CF")
    ws.cell(row=row, column=2).value = "$"
    # Y0 = -initial equity, Y1-Y19 = annual equity CF, Y20 = annual + residual
    apply_calc_style(ws, row, 4, f'=-$D${rows["initial_equity"]}')
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 24):  # Y1-Y19
        c = col_letter(col)
        formula = f'={c}{rows["annual_equity_cf"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    # Y20 = annual CF + residual
    apply_calc_style(ws, row, 24, f'=X{rows["annual_equity_cf"]}+X{rows["life_hold_residual"]}')
    ws.cell(row=row, column=24).number_format = CURRENCY_FMT
    row += 2

    # Short Hold
    rows['short_hold_sale_price'] = row
    set_label(ws, row, "Short Hold: Sale Price (Y5)")
    ws.cell(row=row, column=2).value = "$"
    ws.cell(row=row, column=3).value = "PV of Y6-Y20 CFADS + Residual"
    # Sale price = PV of remaining CFADS (Y6-Y20) + PV of residual at buyer CoE
    # Year 5 = column I (col 9), Year 6 = J, ..., Year 20 = X
    # NPV formula: Sum of CFADS/(1+r)^n for n=1 to 15, plus residual/(1+r)^15
    # Using NPV function: =NPV(rate, J{cfads}:X{cfads}) + residual/(1+rate)^15
    apply_calc_style(ws, row, 9, f'=NPV($D${rows["buyer_coe"]},J{rows["cfads"]}:X{rows["cfads"]})+$D${rows["residual_value"]}/(1+$D${rows["buyer_coe"]})^15')
    ws.cell(row=row, column=9).number_format = CURRENCY_FMT
    ws.cell(row=row, column=9).fill = OUTPUT_FILL
    row += 1

    # Short Hold Equity CF
    rows['short_hold_equity_cf'] = row
    set_label(ws, row, "Short Hold: Equity CF")
    ws.cell(row=row, column=2).value = "$"
    # Y0 = -initial equity, Y1-Y4 = annual equity CF, Y5 = annual + sale price - debt payoff
    apply_calc_style(ws, row, 4, f'=-$D${rows["initial_equity"]}')
    ws.cell(row=row, column=4).number_format = CURRENCY_FMT
    for col in range(5, 9):  # Y1-Y4
        c = col_letter(col)
        formula = f'={c}{rows["annual_equity_cf"]}'
        apply_calc_style(ws, row, col, formula)
        ws.cell(row=row, column=col).number_format = CURRENCY_FMT
    # Y5 = annual CF + sale price - remaining debt
    apply_calc_style(ws, row, 9, f'=I{rows["annual_equity_cf"]}+I{rows["short_hold_sale_price"]}-I{rows["tr1_end"]}-I{rows["tr2_end"]}')
    ws.cell(row=row, column=9).number_format = CURRENCY_FMT
    row += 2

    # ============================================================
    # SECTION 12: RETURNS
    # ============================================================
    set_header(ws, row, "═══ RETURNS ═══")
    row += 1

    # Life Hold Returns
    rows['life_hold_irr'] = row
    set_label(ws, row, "Life Hold: Equity IRR")
    ws.cell(row=row, column=2).value = "%"
    apply_output_style(ws, row, 4, f'=IRR(D{rows["life_hold_equity_cf"]}:X{rows["life_hold_equity_cf"]})')
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['life_hold_coc'] = row
    set_label(ws, row, "Life Hold: Avg Cash-on-Cash")
    ws.cell(row=row, column=2).value = "%"
    apply_output_style(ws, row, 4, f'=AVERAGE(E{rows["annual_equity_cf"]}:X{rows["annual_equity_cf"]})/$D${rows["initial_equity"]}')
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['life_hold_multiple'] = row
    set_label(ws, row, "Life Hold: Whole Dollar Multiple")
    ws.cell(row=row, column=2).value = "x"
    apply_output_style(ws, row, 4, f'=(SUM(E{rows["life_hold_equity_cf"]}:X{rows["life_hold_equity_cf"]}))/-D{rows["life_hold_equity_cf"]}')
    ws.cell(row=row, column=4).number_format = MULTIPLE_FMT
    row += 2

    # Short Hold Returns
    rows['short_hold_irr'] = row
    set_label(ws, row, "Short Hold: Equity IRR")
    ws.cell(row=row, column=2).value = "%"
    apply_output_style(ws, row, 4, f'=IRR(D{rows["short_hold_equity_cf"]}:I{rows["short_hold_equity_cf"]})')
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['short_hold_coc'] = row
    set_label(ws, row, "Short Hold: Avg Cash-on-Cash")
    ws.cell(row=row, column=2).value = "%"
    apply_output_style(ws, row, 4, f'=AVERAGE(E{rows["annual_equity_cf"]}:I{rows["annual_equity_cf"]})/$D${rows["initial_equity"]}')
    ws.cell(row=row, column=4).number_format = PERCENT_FMT
    row += 1

    rows['short_hold_multiple'] = row
    set_label(ws, row, "Short Hold: Whole Dollar Multiple")
    ws.cell(row=row, column=2).value = "x"
    apply_output_style(ws, row, 4, f'=(SUM(E{rows["short_hold_equity_cf"]}:I{rows["short_hold_equity_cf"]}))/-D{rows["short_hold_equity_cf"]}')
    ws.cell(row=row, column=4).number_format = MULTIPLE_FMT
    row += 2

    # ============================================================
    # SECTION 13: SENSITIVITIES (SHORT HOLD ONLY)
    # ============================================================
    set_header(ws, row, "═══ SENSITIVITIES (SHORT HOLD) ═══")
    row += 1

    ws.cell(row=row, column=1).value = "Note: Sensitivities require manual input changes or data tables in Excel"
    ws.cell(row=row, column=1).font = Font(italic=True)
    row += 2

    # EAF Sensitivity table
    rows['sens_eaf_header'] = row
    set_label(ws, row, "EAF Sensitivity")
    ws.cell(row=row, column=1).font = BOLD_FONT
    row += 1

    eaf_values = [0.90, 0.91, 0.92, 0.93, 0.94, 0.95, 0.96, 0.97, 0.98, 0.985, 0.99, 1.00]
    set_label(ws, row, "EAF")
    for i, eaf in enumerate(eaf_values):
        ws.cell(row=row, column=4+i).value = eaf
        ws.cell(row=row, column=4+i).number_format = PERCENT_FMT
    row += 1

    rows['sens_eaf_irr'] = row
    set_label(ws, row, "IRR (base case at EAF=98.5%)")
    ws.cell(row=row, column=3).value = "→"
    apply_output_style(ws, row, 4+9, f'=$D${rows["short_hold_irr"]}')  # 98.5% is at index 9
    ws.cell(row=row, column=4+9).number_format = PERCENT_FMT
    row += 2

    # LIBOR Sensitivity
    rows['sens_libor_header'] = row
    set_label(ws, row, "LIBOR Sensitivity")
    ws.cell(row=row, column=1).font = BOLD_FONT
    row += 1

    set_label(ws, row, "LIBOR")
    ws.cell(row=row, column=4).value = 0.02
    ws.cell(row=row, column=5).value = 0.03
    ws.cell(row=row, column=6).value = 0.04
    for col in range(4, 7):
        ws.cell(row=row, column=col).number_format = PERCENT_FMT
    row += 1

    rows['sens_libor_irr'] = row
    set_label(ws, row, "IRR (base case at 3%)")
    ws.cell(row=row, column=3).value = "→"
    apply_output_style(ws, row, 5, f'=$D${rows["short_hold_irr"]}')
    ws.cell(row=row, column=5).number_format = PERCENT_FMT
    row += 2

    # Buyer CoE Sensitivity
    rows['sens_coe_header'] = row
    set_label(ws, row, "Buyer CoE Sensitivity")
    ws.cell(row=row, column=1).font = BOLD_FONT
    row += 1

    set_label(ws, row, "Buyer CoE")
    ws.cell(row=row, column=4).value = 0.09
    ws.cell(row=row, column=5).value = 0.10
    ws.cell(row=row, column=6).value = 0.11
    ws.cell(row=row, column=7).value = 0.12
    for col in range(4, 8):
        ws.cell(row=row, column=col).number_format = PERCENT_FMT
    row += 1

    rows['sens_coe_irr'] = row
    set_label(ws, row, "IRR (base case at 10%)")
    ws.cell(row=row, column=3).value = "→"
    apply_output_style(ws, row, 5, f'=$D${rows["short_hold_irr"]}')
    ws.cell(row=row, column=5).number_format = PERCENT_FMT
    row += 2

    # ============================================================
    # SECTION 14: QA CHECKS
    # ============================================================
    set_header(ws, row, "═══ QA CHECKS ═══")
    row += 1

    rows['qa_senior_zero'] = row
    set_label(ws, row, "Senior ends at $0 by 2033")
    apply_output_style(ws, row, 4, f'=IF(X{rows["tr1_end"]}<1,"PASS","FAIL")')
    row += 1

    rows['qa_sub_zero'] = row
    set_label(ws, row, "Sub ends at $0 by 2033")
    apply_output_style(ws, row, 4, f'=IF(X{rows["tr2_end"]}<1,"PASS","FAIL")')
    row += 1

    rows['qa_min_dscr'] = row
    set_label(ws, row, "MIN(DSCR) >= 1.40")
    # Use ROUND to avoid floating point precision issues at boundary
    apply_output_style(ws, row, 4, f'=IF(ROUND($D${rows["min_dscr"]},4)>=1.4,"PASS","FAIL")')
    row += 1

    rows['qa_aaf'] = row
    set_label(ws, row, "With EAF=98.5%, AAF=1.00")
    # Use ROUND to avoid floating point precision issues
    apply_output_style(ws, row, 4, f'=IF(ROUND(E{rows["aaf"]},4)=1,"PASS","FAIL")')
    row += 1

    rows['qa_starts'] = row
    set_label(ws, row, "Annual starts (60) <= 1,095")
    apply_output_style(ws, row, 4, f'=IF($D${rows["annual_starts"]}<=1095,"PASS","FAIL")')
    row += 1

    # Save
    output_path = '/mnt/user-data/outputs/Enterprise_Peaker_Model.xlsx'
    wb.save(output_path)
    print(f"Model saved to: {output_path}")

    # Copy to deliverables
    import shutil
    deliv_path = '/home/user/dcf-modeling/deliverables/Enterprise_Peaker_Model.xlsx'
    shutil.copy(output_path, deliv_path)
    print(f"Also saved to: {deliv_path}")

    return rows


if __name__ == "__main__":
    rows = build_model()
    print("\n✓ Enterprise Peaker Model built successfully")
    print(f"  Total rows: {max(rows.values())}")
